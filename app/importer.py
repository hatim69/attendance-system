"""Bulk import from CSV / Excel (.xlsx) so data exported from another system can be loaded in one upload.

Headers are matched case-insensitively against a list of aliases (e.g. "Emp Code", "EmpCode", "Card No" all mean the
employee id), so most exports load without editing. Every import supports a dry run that reports what would happen."""
import csv
import io
import re
from datetime import date, datetime, time as dtime

from . import db, leavepolicy, profile
from .reports import LEAVE_TYPES, WEEKDAYS

MAX_ROWS = 100_000

ALIASES = {
    "id": ["id", "employee id", "emp id", "emp code", "empcode", "employee code", "code", "card no", "cardno", "card no.",
           "enroll no", "enroll no.", "enrollno", "user id", "userid", "device_user_id", "pin", "emp.code"],
    "name": ["name", "employee name", "empname", "emp name", "emp_name", "full name"],
    "shift_name": ["shift name", "name", "shiftname", "shift_name", "shift"],
    "holiday_name": ["holiday name", "name", "holiday", "description", "occasion"],
    "department": ["department", "dept", "dept name", "dept_name", "department name"],
    "shift": ["shift", "shift name", "shiftname"],
    "designation": ["designation", "desig name", "desig_name", "desig", "designation name"],
    "branch": ["branch", "branch name", "branch_name", "location"],
    "mobile": ["mobile", "mobile number", "mobile no.", "phone", "phone.no", "phone no", "phone no.", "mobile no", "contact", "contact no"],
    "father": ["father name", "father's name", "father/husband name", "husband name", "fathername",
               "father's name, occupation and contact number"],
    "salary": ["salary", "monthly salary", "gross salary", "basic"],
    "joined": ["joined", "date of joining", "dateof join", "date of join", "doj", "joining date", "join date"],
    "start": ["start", "start time", "shift start", "from time", "in time"],
    "end": ["end", "end time", "shift end", "to time", "out time"],
    "grace": ["grace", "grace min", "grace_min", "grace minutes", "late grace"],
    "half_day_hours": ["half day hours", "half_day_hours", "half day hrs", "half-day hours"],
    "ot_after": ["ot after", "ot_after_min", "overtime after", "ot after min"],
    "weekly_off": ["weekly off", "weekly_off", "weekoff", "week off", "weekly offs"],
    "date": ["date", "holiday date", "day", "punch date", "attendance date"],
    "from": ["from", "from date", "start date", "leave from", "leave from date", "date from"],
    "to": ["to", "to date", "end date", "leave to", "leave to date", "date to"],
    "type": ["type", "leave type", "leave name", "leave"],
    "note": ["note", "reason", "remark", "remarks", "comment"],
    "datetime": ["datetime", "date time", "attendance datetime", "punch time", "timestamp", "log time", "attendancedatetime"],
    "time": ["time", "times", "swipe record", "punches", "swipes"],
    "device": ["device", "device id", "deviceid", "device name", "serial", "machine"],
    # personal and bank details (the old system's employee export and the MTT employee information form)
    "dob": ["dob", "date of birth", "birth date", "birthday"],
    "gender": ["gender", "sex"],
    "blood_group": ["blood group", "blood"],
    "nationality": ["nationality"],
    "email": ["personal email address", "personal email", "email", "email address", "e-mail", "email id"],
    "work_email": ["company mail address", "company email", "work email", "office email", "official email"],
    "address": ["address", "house/hostel address", "residential address", "current address", "permanent address"],
    "alt_contact": ["alternate number, name and relation", "alternate number", "emergency contact", "alternate contact"],
    "mother_name": ["mother's name occupation and contact number", "mother's name", "mother name"],
    "spouse_name": ["spouse name occupation and contact number", "wife's name occupation and contact number",
                    "spouse name", "spouse", "wife's name", "husband's name"],
    "govt_uid": ["govt. uid", "govt uid", "aadhaar", "aadhar", "aadhaar number", "aadhar number", "aadhaar no", "aadhar no", "uid"],
    "pan": ["pan number", "pan", "pan no", "pan no.", "pan card number"],
    "passport_no": ["passport number", "passport no", "passport no.", "passport"],
    "passport_expiry": ["passport expiry date", "passport expiry", "passport valid till"],
    "bank_name": ["bank name", "bank"],
    "bank_account": ["bank account number", "bank a/c number", "account number", "account no", "a/c no", "bank account", "a/c number"],
    "bank_ifsc": ["ifsc code", "bank ifsc code", "ifsc", "bank ifsc"],
    "bank_holder": ["account holders name", "account holder's name", "account holder name", "account holder"],
    "bank_type": ["account type", "bank account type", "type of account"],
}
TEMPLATES = {
    "employees": ["id", "name", "department", "designation", "branch", "mobile", "shift", "salary", "joined"],
    "shifts": ["name", "start", "end", "grace", "half_day_hours", "ot_after", "weekly_off"],
    "holidays": ["date", "name"],
    "leaves": ["id", "from", "to", "type", "note"],
    "punches": ["id", "datetime", "device"],
    "profiles": ["id", "dob", "gender", "blood_group", "address", "mobile", "email", "govt_uid", "pan",
                 "bank_name", "bank_account", "bank_ifsc", "bank_holder", "bank_type"],
}
# HR-kept fields that a details sheet only fills in when they are still empty (it never overwrites them)
PROFILE_FILL = {"mobile": "mobile", "father": "father_name", "joined": "joined"}
TEMPLATE_EXAMPLES = {
    "employees": [["217315", "Asha Khan", "Sales", "Executive", "Mumbai", "9876543210", "Morning 9-6", "30000", "2025-04-01"]],
    "shifts": [["Morning 9-6", "09:00", "18:00", "10", "4", "30", "SUN"]],
    "holidays": [["2026-10-20", "Dussehra"]],
    "leaves": [["217315", "2026-10-12", "2026-10-13", "Informed", "family function"]],
    "punches": [["217315", "2026-10-05 09:12:00", "Mumbai Office"]],
    "profiles": [["217315", "1995-06-21", "Female", "B+", "12 MG Road, Secunderabad", "9876543210", "asha@example.com",
                  "123412341234", "ABCDE1234F", "State Bank of India", "12345678901", "SBIN0001234", "Asha Khan", "Savings"]],
}


class ImportError_(Exception):
    pass


def _norm(h) -> str:
    """Header key: lower case, single spaces, without notes in brackets ("Address (if in hostel, both)" -> "address")."""
    return re.sub(r"\s+", " ", re.sub(r"\(.*?\)", " ", str(h or ""))).strip().lower()


def to_id(v) -> str:
    """Employee id from a cell: Excel stores 217204 as the number 217204.0."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return db.canon_id(v)


def read_table(filename: str, data: bytes) -> list[dict]:
    """Rows as dicts keyed by the normalised header. Supports .csv and .xlsx."""
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    elif name.endswith((".csv", ".txt")):
        text = data.decode("utf-8-sig", "replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    else:
        raise ImportError_("Upload a .csv or .xlsx file")
    rows = [r for r in rows if any(str(c or "").strip() for c in r)]
    if len(rows) < 2:
        raise ImportError_("The file has no data rows")
    if len(rows) > MAX_ROWS + 1:
        raise ImportError_(f"Too many rows (max {MAX_ROWS})")
    head = [_norm(h) for h in rows[0]]
    return [{h: (r[i] if i < len(r) else None) for i, h in enumerate(head) if h} for r in rows[1:]]


def pick(row: dict, field: str):
    for a in ALIASES[field]:
        v = row.get(a)
        if v is not None and str(v).strip() != "":
            return v
    return None


def to_date(v) -> str | None:
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    t = str(v).strip().split(" ")[0]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(t, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"unreadable date {v!r} (use YYYY-MM-DD or DD/MM/YYYY)")


def to_hm(v) -> str:
    if isinstance(v, datetime):
        return v.strftime("%H:%M")
    if isinstance(v, dtime):
        return v.strftime("%H:%M")
    m = re.fullmatch(r"\s*(\d{1,2})[:.](\d{2})(?::\d{2})?\s*([AaPp][Mm])?\s*", str(v))
    if not m:
        raise ValueError(f"unreadable time {v!r} (use HH:MM)")
    h, mi = int(m.group(1)), int(m.group(2))
    if m.group(3):
        h = h % 12 + (12 if m.group(3).lower() == "pm" else 0)
    if h > 23 or mi > 59:
        raise ValueError(f"unreadable time {v!r}")
    return f"{h:02d}:{mi:02d}"


def to_dt(v) -> str:
    """'YYYY-MM-DD HH:MM[:SS]' (also DD/MM/YYYY and a 'T' separator) -> 'YYYY-MM-DD HH:MM:SS'."""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    d, _, t = str(v).strip().replace("T", " ").partition(" ")
    t = t.strip()
    if not t:
        raise ValueError(f"missing time in {v!r}")
    sec = re.match(r"\d{1,2}[:.]\d{2}[:.](\d{2})", t)
    return f"{to_date(d)} {to_hm(t)}:{sec.group(1) if sec else '00'}"


def to_num(v, default=0.0) -> float:
    if v is None or str(v).strip() == "":
        return default
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        raise ValueError(f"not a number: {v!r}")


def run(kind: str, filename: str, data: bytes, dry: bool, inactive: bool = False) -> dict:
    if kind not in TEMPLATES:
        raise ImportError_("Unknown import type")
    rows = read_table(filename, data)
    c = db.conn()
    res = {"kind": kind, "dry_run": dry, "rows": len(rows), "created": 0, "updated": 0, "skipped": 0, "errors": []}
    shifts = {r["name"].lower(): r["id"] for r in c.execute("SELECT id, name FROM shifts")}
    known_ids = {r[0] for r in c.execute("SELECT device_user_id FROM employees")}

    def err(i, msg):
        res["errors"].append({"row": i + 2, "message": msg})  # +2: header is row 1
        res["skipped"] += 1

    for i, row in enumerate(rows):
        try:
            if kind == "employees":
                uid, name = to_id(pick(row, "id")), str(pick(row, "name") or "").strip()
                if not uid or not name:
                    raise ValueError("needs an id and a name")
                sh = pick(row, "shift")
                sid = shifts.get(str(sh).strip().lower()) if sh else None
                if sh and sid is None:
                    raise ValueError(f"unknown shift {sh!r} (import shifts first)")
                text = lambda f: profile.text(pick(row, f))
                vals = {"name": name, "department": text("department"), "designation": text("designation"),
                        "branch": text("branch"), "mobile": text("mobile"), "father_name": text("father"),
                        "salary": to_num(pick(row, "salary")), "joined": to_date(pick(row, "joined")), "shift_id": sid}
                vals.update(profile.clean_all({f: pick(row, f) for f in db.PROFILE_FIELDS}, strict=False))
                if inactive:       # the file lists resigned staff: keep their history but not in active reports
                    vals["active"] = 0
                exists = uid in known_ids
                if not dry:
                    if exists:   # only overwrite what the file provides; never touches PINs
                        sets = {k: v for k, v in vals.items() if (v not in (None, "") and not (k == "salary" and v == 0)) or k == "active"}
                        c.execute("UPDATE employees SET " + ",".join(f"{k}=?" for k in sets) + " WHERE device_user_id=?",
                                  list(sets.values()) + [uid])
                    else:
                        vals.setdefault("active", 1)
                        c.execute(f"INSERT INTO employees(device_user_id,{','.join(vals)}) VALUES (?{',?' * len(vals)})",
                                  [uid] + list(vals.values()))
                known_ids.add(uid)
                res["updated" if exists else "created"] += 1
            elif kind == "profiles":   # personal / bank details for people who already exist; work details untouched
                uid = to_id(pick(row, "id"))
                if not uid:
                    raise ValueError("needs an employee id")
                if uid not in known_ids:
                    raise ValueError(f"no employee with id {uid} (add them in Employees first)")
                vals = {k: v for k, v in profile.clean_all({f: pick(row, f) for f in db.PROFILE_FIELDS}, strict=False).items() if v}
                cur = c.execute("SELECT mobile, father_name, COALESCE(joined,'') joined FROM employees WHERE device_user_id=?", (uid,)).fetchone()
                for alias, col in PROFILE_FILL.items():
                    v = pick(row, alias)
                    if v is not None and not cur[col]:
                        vals[col] = to_date(v) if col == "joined" else profile.text(v)
                vals = {k: v for k, v in vals.items() if v}
                if not vals:
                    res["skipped"] += 1
                    continue
                if not dry:
                    c.execute("UPDATE employees SET " + ",".join(f"{k}=?" for k in vals) + " WHERE device_user_id=?",
                              list(vals.values()) + [uid])
                res["updated"] += 1
            elif kind == "shifts":
                name = str(pick(row, "shift_name") or "").strip()
                start, end = to_hm(pick(row, "start")), to_hm(pick(row, "end"))
                if not name:
                    raise ValueError("needs a shift name")
                if end <= start:
                    raise ValueError("end must be after start (overnight shifts are not supported yet)")
                offs = [o.strip().upper()[:3] for o in re.split(r"[,/; ]+", str(pick(row, "weekly_off") or "SUN")) if o.strip()]
                if any(o not in WEEKDAYS for o in offs):
                    raise ValueError("weekly off must be like SUN or SAT,SUN")
                vals = (start, end, int(to_num(pick(row, "grace"), 10)), to_num(pick(row, "half_day_hours"), 4),
                        int(to_num(pick(row, "ot_after"), 30)), ",".join(offs))
                exists = name.lower() in shifts
                if not dry:
                    if exists:
                        c.execute("UPDATE shifts SET start=?, end=?, grace_min=?, half_day_hours=?, ot_after_min=?, weekly_off=? WHERE id=?",
                                  vals + (shifts[name.lower()],))
                    else:
                        shifts[name.lower()] = c.execute(
                            "INSERT INTO shifts(name,start,end,grace_min,half_day_hours,ot_after_min,weekly_off) VALUES (?,?,?,?,?,?,?)",
                            (name,) + vals).lastrowid
                else:
                    shifts.setdefault(name.lower(), -1)
                res["updated" if exists else "created"] += 1
            elif kind == "holidays":
                d, name = to_date(pick(row, "date")), str(pick(row, "holiday_name") or "").strip()
                if not d or not name:
                    raise ValueError("needs a date and a name")
                exists = c.execute("SELECT 1 FROM holidays WHERE date=?", (d,)).fetchone() is not None
                if not dry:
                    c.execute("INSERT INTO holidays(date,name) VALUES (?,?) ON CONFLICT(date) DO UPDATE SET name=excluded.name", (d, name))
                res["updated" if exists else "created"] += 1
            elif kind == "leaves":
                uid, d1, d2 = to_id(pick(row, "id")), to_date(pick(row, "from")), to_date(pick(row, "to")) or to_date(pick(row, "from"))
                if not uid or not d1:
                    raise ValueError("needs an id and a from date")
                if d2 < d1:
                    raise ValueError("to date is before from date")
                raw = str(pick(row, "type") or "Informed").strip()
                ltype = leavepolicy.ALIASES.get(raw.upper() if raw.upper() == "CL" else raw.title(), raw.title())
                if ltype not in LEAVE_TYPES:
                    raise ValueError(f"leave type {raw!r} must be Informed or Adhoc (CL / Casual = Informed, Sick = Adhoc)")
                if uid not in known_ids:
                    raise ValueError(f"no employee with id {uid} (import employees first)")
                dup = c.execute("SELECT 1 FROM leaves WHERE device_user_id=? AND start_date=? AND end_date=?", (uid, d1, d2)).fetchone()
                if dup:
                    res["skipped"] += 1
                    continue
                if not dry:
                    c.execute("INSERT INTO leaves(device_user_id,start_date,end_date,type,note) VALUES (?,?,?,?,?)",
                              (uid, d1, d2, ltype, str(pick(row, "note") or "")))
                res["created"] += 1
            else:  # punches
                uid = to_id(pick(row, "id"))
                if not uid:
                    raise ValueError("needs an employee id")
                source = str(pick(row, "device") or "import").strip()
                stamps = []
                if pick(row, "datetime") is not None:
                    stamps = [to_dt(pick(row, "datetime"))]
                elif pick(row, "date") is not None and pick(row, "time") is not None:
                    d = to_date(pick(row, "date"))   # one or several times in a cell: "09:58 13:00 18:30"
                    tv = pick(row, "time")
                    parts = [tv] if isinstance(tv, (datetime, dtime)) else re.findall(r"\d{1,2}[:.]\d{2}(?::\d{2})?\s*(?:[AaPp][Mm])?", str(tv))
                    stamps = [f"{d} {to_hm(p)}:00" for p in parts]
                if not stamps:
                    raise ValueError("needs a datetime, or a date plus time(s)")
                for ts in stamps:
                    stored = None if dry else db.add_punch(uid, ts, 0, source, method="import")
                    if dry:
                        stored = not c.execute("SELECT 1 FROM punches WHERE device_user_id=? AND ts=?", (uid, ts)).fetchone()
                    res["created" if stored else "skipped"] += 1
        except ValueError as e:
            err(i, str(e))
    if not dry:
        c.commit()
    else:
        c.rollback()
    res["errors"] = res["errors"][:200]
    return res


def template_csv(kind: str) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(TEMPLATES[kind])
    w.writerows(TEMPLATE_EXAMPLES[kind])
    return buf.getvalue()
