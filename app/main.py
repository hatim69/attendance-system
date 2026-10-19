import asyncio
import csv
import html
import io
import logging
import secrets
from contextlib import asynccontextmanager
from datetime import date, datetime

import os
import re

from fastapi import Depends, FastAPI, File, HTTPException, Request, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, PlainTextResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from .timeutil import now_local

from .xlsx import TIMECARD_COLUMNS, XLSX_TYPE, inout_sheet, safe as _safe, timecard_human, xlsx_sheet
from . import access, adms, app_api, auth, backup, config, db, fkweb, hr_api, hub, importer, leavepolicy, loans, lates, mailer, payslip, profile, push, reports, weekly, zk_poller

logging.basicConfig(level=logging.INFO)


@asynccontextmanager
async def lifespan(app: FastAPI):
    db.init()
    access.bootstrap()
    if access.open_mode() and config.REQUIRE_LOGIN:
        raise RuntimeError("REQUIRE_LOGIN is set but there are no users and ADMIN_PASSWORD is empty. Set ADMIN_PASSWORD in .env.")
    if access.open_mode():
        logging.warning("OPEN MODE: no users and no ADMIN_PASSWORD, so the dashboard needs no login. Set ADMIN_PASSWORD before hosting!")
    hub.bind_loop(asyncio.get_running_loop())
    zk_poller.start()
    fk_server = await fkweb.start()
    reminders = asyncio.create_task(push.loop()) if config.PUSH_REMINDERS else None
    nightly = asyncio.create_task(backup.loop()) if config.BACKUP_HOUR >= 0 else None
    late_mail = asyncio.create_task(lates.loop()) if config.LATE_EMAILS else None
    weekly_mail = asyncio.create_task(weekly.loop()) if config.WEEKLY_REPORT else None
    yield
    for t in (reminders, nightly, late_mail, weekly_mail):
        if t:
            t.cancel()
    if fk_server:
        fk_server.close()


app = FastAPI(title="Attendance System", lifespan=lifespan)
app.include_router(adms.router)  # device-facing, no auth (devices can't do Basic auth)
app.include_router(app_api.router)  # employee app, token auth
app.include_router(hr_api.router)   # documents, pick lists, loans, roster



class EmployeeIn(BaseModel):
    device_user_id: str
    name: str
    department: str = ""
    active: bool = True
    shift_id: int | None = None
    salary: float | None = None   # left out = unchanged (team leads never see salaries); 0 for someone new
    joined: str = ""  # optional YYYY-MM-DD; days before it are not counted
    designation: str = ""
    branch: str = ""
    mobile: str = ""
    father_name: str | None = None
    remote_ok: bool = False   # may check in with the phone app from anywhere (work from home / field work)
    left_on: str = ""         # last working day, for someone who resigned
    # personal and bank details (db.PROFILE_FIELDS); left out = unchanged on an existing employee
    dob: str | None = None
    gender: str | None = None
    blood_group: str | None = None
    nationality: str | None = None
    email: str | None = None
    work_email: str | None = None
    address: str | None = None
    alt_contact: str | None = None
    mother_name: str | None = None
    spouse_name: str | None = None
    govt_uid: str | None = None
    pan: str | None = None
    passport_no: str | None = None
    passport_expiry: str | None = None
    bank_name: str | None = None
    bank_account: str | None = None
    bank_ifsc: str | None = None
    bank_holder: str | None = None
    bank_type: str | None = None


class ShiftIn(BaseModel):
    id: int | None = None
    name: str
    start: str
    end: str
    grace_min: int = 10
    half_day_hours: float = 4
    ot_after_min: int = 30
    weekly_off: str = "SUN"  # comma list of MON..SUN, empty for none
    break_start: str = ""    # optional unpaid break, HH:MM
    break_end: str = ""


class HolidayIn(BaseModel):
    date: str
    name: str
    optional: bool = False   # no longer used (the festival list was removed); kept so older clients still work


class LeaveIn(BaseModel):
    device_user_id: str
    start_date: str
    end_date: str
    type: str = "Informed"
    note: str = ""
    applied_on: str = ""      # when the application was made (default today); matters for Normal leave's 7-day notice
    override: bool = False    # record exactly as given, skipping the leave policy (recorded in the audit log)
    half: bool = False        # half-day leave (one date)


class DeviceIn(BaseModel):
    name: str = ""
    location: str = ""


class PinIn(BaseModel):
    pin: str


class SiteIn(BaseModel):
    name: str
    lat: float
    lng: float
    radius_m: float = 150
    kind: str = "office"            # office | home (one employee's work-from-home location)
    device_user_id: str = ""        # for home sites


class ReviewIn(BaseModel):
    flag: str  # 'rejected' | 'approved' | ''


class ShiftChangeIn(BaseModel):
    ids: list[str]
    start_date: str
    end_date: str
    shift_id: int


class DayPunchIn(BaseModel):
    employee_id: str
    date: str
    in_time: str = ""
    out_time: str = ""
    remark: str = ""


class CorrectIn(BaseModel):
    ts: str
    remark: str = ""


class PunchIn(BaseModel):
    device_user_id: str
    ts: str  # "YYYY-MM-DD HH:MM:SS"
    status: int = 0


class LoginIn(BaseModel):
    username: str
    password: str


class PasswordIn(BaseModel):
    old_password: str
    new_password: str


class UserIn(BaseModel):
    username: str
    display_name: str = ""
    password: str = ""
    role: str = "manager"
    department: str = ""
    active: bool = True
    permissions: str = "*"   # hr only: comma list of areas (employees,shifts,leave,holidays,punches,import,payroll) or *
    email: str = ""          # team leads receive late-arrival e-mails for their team here


def _me(u: dict) -> dict:
    return {"username": u["username"], "role": u["role"], "department": u["department"],
            "display_name": u["display_name"], "open_mode": u["username"] == "(open)",
            "permissions": u.get("permissions", "*"), "company": config.COMPANY_NAME}


@app.post("/api/auth/login")
def auth_login(body: LoginIn, request: Request):
    name = body.username.strip().lower()
    if not name or not body.password:
        raise HTTPException(400, "Type your username and password")
    key = "u:" + name
    if auth.locked(key):
        raise HTTPException(429, "Too many failed attempts. Try again in 5 minutes.")
    row = db.conn().execute("SELECT * FROM users WHERE username=? AND active=1", (name,)).fetchone()
    if not row or not auth.check_password(body.password, row["password_hash"]):
        auth.record_fail(key)
        db.audit(name, "LOGIN failed", 401)
        raise HTTPException(401, "Wrong username or password")
    auth.record_ok(key)
    db.audit(name, "LOGIN", 200)
    resp = JSONResponse(_me(dict(row)))
    resp.set_cookie(access.COOKIE, auth.make_token(name, "admin", access.SESSION_TTL), max_age=access.SESSION_TTL,
                    httponly=True, samesite="strict", secure=request.url.scheme == "https", path="/")
    return resp


@app.post("/api/auth/logout")
def auth_logout():
    resp = JSONResponse({"ok": True})
    resp.delete_cookie(access.COOKIE, path="/")
    return resp


@app.get("/api/auth/me")
async def auth_me(u: dict = Depends(access.need("manager"))):
    return _me(u)


@app.post("/api/auth/password")
async def auth_password(body: PasswordIn, u: dict = Depends(access.need("manager"))):
    if u["username"] == "(open)":
        raise HTTPException(400, "Create a user first (set ADMIN_PASSWORD and restart)")
    if len(body.new_password) < 8:
        raise HTTPException(400, "New password must be at least 8 characters")
    if not auth.check_password(body.old_password, u["password_hash"]):
        raise HTTPException(403, "Current password is wrong")
    c = db.conn()
    c.execute("UPDATE users SET password_hash=? WHERE username=?", (auth.hash_password(body.new_password), u["username"]))
    c.commit()
    return {"ok": True}


@app.get("/api/users", dependencies=[Depends(access.need("admin"))])
def list_users():
    return [dict(r) for r in db.conn().execute("SELECT username, display_name, role, department, active, permissions, email FROM users ORDER BY username")]


@app.post("/api/users")
async def save_user(body: UserIn, u: dict = Depends(access.need("admin"))):
    name = body.username.strip().lower()
    if not re.fullmatch(r"[a-z0-9._-]{3,30}", name):
        raise HTTPException(400, "Username: 3-30 characters, letters, digits . _ -")
    if body.role not in access.LEVEL:
        raise HTTPException(400, "Role must be admin, hr, management or manager")
    if body.role == "manager" and not body.department.strip():
        raise HTTPException(400, "A team lead needs a department")
    perms = [x.strip() for x in body.permissions.split(",") if x.strip()]
    if body.role == "manager":
        perms = [x for x in perms if x != "*"]      # "*" only means "the HR defaults"
    if body.role == "hr":
        if perms != ["*"] and any(x not in access.ALL_PERMS for x in perms):
            raise HTTPException(400, "Abilities must be * or a list of: " + ", ".join(access.ALL_PERMS))
        perm_text = "*" if perms in ([], ["*"]) else ",".join(perms)
    elif body.role == "manager":
        if any(x not in access.MANAGER_GRANTABLE for x in perms):
            raise HTTPException(400, "A team lead can be granted only: " + ", ".join(access.MANAGER_GRANTABLE))
        perm_text = ",".join(perms)
    else:
        perm_text = "*"      # admin and management always have everything
    if body.password and len(body.password) < 8:
        raise HTTPException(400, "Password must be at least 8 characters")
    try:
        email = profile.clean("email", body.email, strict=True)
    except ValueError as ex:
        raise HTTPException(400, str(ex))
    c = db.conn()
    row = c.execute("SELECT * FROM users WHERE username=?", (name,)).fetchone()
    if not row and not body.password:
        raise HTTPException(400, "A password is required for a new user")
    if row and row["role"] == "admin" and (body.role != "admin" or not body.active):
        if c.execute("SELECT COUNT(*) FROM users WHERE role='admin' AND active=1 AND username!=?", (name,)).fetchone()[0] == 0:
            raise HTTPException(400, "There must be at least one active admin")
    if row:
        c.execute("UPDATE users SET display_name=?, role=?, department=?, active=?, permissions=?, email=? WHERE username=?",
                  (body.display_name, body.role, body.department.strip(), int(body.active), perm_text, email, name))
        if body.password:
            c.execute("UPDATE users SET password_hash=? WHERE username=?", (auth.hash_password(body.password), name))
    else:
        c.execute("INSERT INTO users(username,display_name,password_hash,role,department,active,permissions,email) VALUES (?,?,?,?,?,?,?,?)",
                  (name, body.display_name, auth.hash_password(body.password), body.role, body.department.strip(), int(body.active), perm_text, email))
    c.commit()
    return {"ok": True}


@app.delete("/api/users/{username}")
async def delete_user(username: str, u: dict = Depends(access.need("admin"))):
    username = username.lower()
    c = db.conn()
    row = c.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
    if not row:
        raise HTTPException(404, "No such user")
    if username == u["username"]:
        raise HTTPException(400, "You cannot delete your own account")
    if row["role"] == "admin" and c.execute("SELECT COUNT(*) FROM users WHERE role='admin' AND active=1 AND username!=?", (username,)).fetchone()[0] == 0:
        raise HTTPException(400, "There must be at least one active admin")
    c.execute("DELETE FROM users WHERE username=?", (username,))
    c.commit()
    return {"ok": True}


@app.get("/api/audit", dependencies=[Depends(access.need("admin"))])
def list_audit(limit: int = 300, date_from: str = "", date_to: str = "", q: str = ""):
    """Newest first, optionally between two dates (inclusive) and matching a user / action text."""
    where, args = [], []
    for d in (date_from, date_to):
        if d:
            _valid_date(d)
    if date_from:
        where, args = where + ["ts >= ?"], args + [date_from]
    if date_to:
        where, args = where + ["ts < date(?, '+1 day')"], args + [date_to]
    if q.strip():
        where, args = where + ["(user LIKE ? OR action LIKE ?)"], args + [f"%{q.strip()}%"] * 2
    sql = "SELECT * FROM audit" + (" WHERE " + " AND ".join(where) if where else "") + " ORDER BY id DESC LIMIT ?"
    return [dict(r) for r in db.conn().execute(sql, (*args, max(1, min(limit, 5000))))]


@app.middleware("http")
async def audit_writes(request: Request, call_next):
    """Record who changed what (method + path, never request bodies, so PINs and passwords are not logged)."""
    response = await call_next(request)
    path = request.url.path
    if request.method not in access.SAFE and path.startswith("/api/") and not path.startswith(("/api/app/", "/api/auth/")):
        db.audit(getattr(request.state, "user", "anonymous"), f"{request.method} {path}", response.status_code)
    return response


@app.get("/api/employees", dependencies=[Depends(access.need("manager"))])
def list_employees():
    frag, args = access.scope_sql("e.department")
    rows = [dict(r) for r in db.conn().execute(
        "SELECT e.id, e.device_user_id, e.name, e.department, e.designation, e.branch, e.mobile, e.father_name, e.active, e.pin_hash IS NOT NULL AS pin_hash, e.shift_id, "
        "e.salary, COALESCE(e.joined,'') joined, e.remote_ok, COALESCE(e.left_on,'') left_on, COALESCE(s.name,'') shift, "
        + "".join(f"e.{f}, " for f in db.PROFILE_FIELDS) +
        "(SELECT COUNT(*) FROM push_subscriptions p WHERE p.device_user_id=e.device_user_id) reminders, "
        "(SELECT COUNT(*) FROM employee_documents d WHERE d.device_user_id=e.device_user_id) documents FROM employees e LEFT JOIN shifts s ON s.id=e.shift_id "
        "WHERE 1=1" + frag + " ORDER BY e.name", args)]
    if not access.salary_visible.get():       # only hr and admin see salaries and personal / bank details
        for r in rows:
            r["salary"] = 0
            for f in db.PROFILE_FIELDS:
                r.pop(f)
    return rows


@app.post("/api/employees", dependencies=[Depends(access.need("hr", "employees"))])
def upsert_employee(e: EmployeeIn):
    e.device_user_id = db.canon_id(e.device_user_id)
    if not e.device_user_id:
        raise HTTPException(400, "Employee ID is required")
    team = access.scope_dept.get()
    if team is not None:   # a team lead with the 'employees' ability stays inside their own team
        existing = db.conn().execute("SELECT department FROM employees WHERE device_user_id=?", (e.device_user_id,)).fetchone()
        if e.department != team or (existing and existing["department"] != team):
            raise HTTPException(403, "You can only add or edit people in your own team")
    for d in (e.joined, e.left_on):
        if d:
            _valid_date(d)
    if e.active:
        e.left_on = ""
    c = db.conn()
    old = c.execute("SELECT * FROM employees WHERE device_user_id=?", (e.device_user_id,)).fetchone()
    # Only the details being changed are checked strictly: an untidy value kept from an import must not stop every
    # other change to that person from saving.
    given = {f: v for f in db.PROFILE_FIELDS if (v := getattr(e, f)) is not None
             and not (old and profile.text(v) == profile.text(old[f]))}
    try:
        prof = profile.clean_all(given, strict=True)
    except ValueError as ex:
        raise HTTPException(400, str(ex))
    if e.salary is None:
        e.salary = old["salary"] if old else 0
    if e.father_name is None:
        e.father_name = old["father_name"] if old else ""
    c.execute(
        "INSERT INTO employees(device_user_id,name,department,active,shift_id,salary,joined,designation,branch,mobile,father_name,remote_ok,left_on) "
        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?) "
        "ON CONFLICT(device_user_id) DO UPDATE SET name=excluded.name, department=excluded.department, "
        "active=excluded.active, shift_id=excluded.shift_id, salary=excluded.salary, joined=excluded.joined, "
        "designation=excluded.designation, branch=excluded.branch, mobile=excluded.mobile, father_name=excluded.father_name, "
        "remote_ok=excluded.remote_ok, left_on=excluded.left_on",
        (e.device_user_id, e.name, e.department, int(e.active), e.shift_id, e.salary, e.joined or None,
         e.designation, e.branch, e.mobile, e.father_name, int(e.remote_ok), e.left_on or None),
    )
    if prof:
        c.execute("UPDATE employees SET " + ",".join(f"{f}=?" for f in prof) + " WHERE device_user_id=?",
                  list(prof.values()) + [e.device_user_id])
    for kind, v in (("department", e.department), ("designation", e.designation)):
        if v.strip():
            c.execute("INSERT OR IGNORE INTO masters(kind,name) VALUES (?,?)", (kind, v.strip()))
    c.commit()
    return {"ok": True}


class ChangeIdIn(BaseModel):
    new_id: str


@app.post("/api/employees/{device_user_id}/change-id", dependencies=[Depends(access.need("hr", "employees"))])
def change_employee_id(device_user_id: str, body: ChangeIdIn):
    """Correct an employee's ID. Everything recorded under the old ID (punches, leave, roster, documents, ...) moves
    with them. The ID must match the one on the machine, so it has to be changed there too."""
    old, new = db.canon_id(device_user_id), db.canon_id(body.new_id)
    access.assert_in_scope(old)
    if not new:
        raise HTTPException(400, "Type the new employee ID")
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,32}", new):
        raise HTTPException(400, "The ID can only have letters, numbers, - and _")
    c = db.conn()
    if not c.execute("SELECT 1 FROM employees WHERE device_user_id=?", (old,)).fetchone():
        raise HTTPException(404, "No such employee")
    if new == old:
        return {"ok": True, "device_user_id": new}
    other = c.execute("SELECT name FROM employees WHERE device_user_id=?", (new,)).fetchone()
    if other:
        raise HTTPException(409, f"ID {new} already belongs to {other['name']}")
    tables = [t["name"] for t in c.execute("SELECT name FROM sqlite_master WHERE type='table'")
              if any(col["name"] == "device_user_id" for col in c.execute(f"PRAGMA table_info({t['name']})"))]
    for t in tables:   # punches the machine already sent under the new ID stay; exact duplicates are dropped
        c.execute(f"UPDATE OR IGNORE {t} SET device_user_id=? WHERE device_user_id=?", (new, old))
        c.execute(f"DELETE FROM {t} WHERE device_user_id=?", (old,))
    c.commit()
    return {"ok": True, "device_user_id": new}


class BulkShiftIn(BaseModel):
    ids: list[str]
    shift_id: int | None = None


@app.post("/api/employees/bulk-shift", dependencies=[Depends(access.need("hr", "employees"))])
def bulk_shift(body: BulkShiftIn):
    """Assign one shift (or none) to many employees at once."""
    c = db.conn()
    if body.shift_id is not None and not c.execute("SELECT 1 FROM shifts WHERE id=?", (body.shift_id,)).fetchone():
        raise HTTPException(404, "No such shift")
    ids = [db.canon_id(i) for i in body.ids]
    if not ids:
        raise HTTPException(400, "Select at least one employee")
    for i in ids:
        access.assert_in_scope(i)
    n = c.executemany("UPDATE employees SET shift_id=? WHERE device_user_id=?", [(body.shift_id, i) for i in ids]).rowcount
    c.commit()
    return {"updated": n}


@app.delete("/api/employees/{device_user_id}", dependencies=[Depends(access.need("hr", "employees"))])
def delete_employee(device_user_id: str):
    device_user_id = db.canon_id(device_user_id)
    access.assert_in_scope(device_user_id)
    c = db.conn()
    c.execute("UPDATE employees SET active=0 WHERE device_user_id=?", (device_user_id,))
    c.commit()
    return {"ok": True}


@app.post("/api/employees/{device_user_id}/pin", dependencies=[Depends(access.need("hr", "employees"))])
def set_pin(device_user_id: str, body: PinIn):
    device_user_id = db.canon_id(device_user_id)
    access.assert_in_scope(device_user_id)
    if not auth.valid_pin(body.pin):
        raise HTTPException(400, "PIN must be 4-8 digits")
    c = db.conn()
    cur = c.execute("UPDATE employees SET pin_hash=? WHERE device_user_id=?", (auth.hash_pin(body.pin), device_user_id))
    c.commit()
    if cur.rowcount == 0:
        raise HTTPException(404, "No such employee")
    return {"ok": True}


ONLINE_SECONDS = 90  # devices poll every ~3 s, so silence for 90 s means offline


@app.get("/api/devices", dependencies=[Depends(access.need("manager"))])
def list_devices():
    now, out = now_local(), []
    for r in db.conn().execute("SELECT * FROM devices ORDER BY COALESCE(NULLIF(name,''), dev_id)"):
        d = dict(r)
        d["online"] = bool(d["last_seen"]) and (now - datetime.fromisoformat(d["last_seen"])).total_seconds() <= ONLINE_SECONDS
        out.append(d)
    return out


@app.post("/api/devices/{dev_id}/info", dependencies=[Depends(access.need("admin", "devices"))])
def update_device(dev_id: str, body: DeviceIn):
    c = db.conn()
    cur = c.execute("UPDATE devices SET name=?, location=? WHERE dev_id=?", (body.name, body.location, dev_id))
    c.commit()
    if cur.rowcount == 0:
        raise HTTPException(404, "Unknown device")
    return {"ok": True}


@app.get("/api/dashboard", dependencies=[Depends(access.need("manager"))])
def dashboard(day: str | None = None, department: str = ""):
    day = day or now_local().date().isoformat()
    _valid_date(day)
    rows = reports.daily(day)
    dept_options = sorted({r["department"] for r in rows if r["department"]})
    if department:
        rows = [r for r in rows if r["department"] == department]
    watch = [x for x in lates.counts(date.fromisoformat(day)) if not department or x["department"] == department]
    depts: dict[str, int] = {}
    for r in rows:
        if r["status"] in reports.WORKED:
            depts[r["department"] or "Unassigned"] = depts.get(r["department"] or "Unassigned", 0) + 1
    c = db.conn()
    stats = {r["source"]: (r["n"], r["last"]) for r in c.execute(
        "SELECT source, COUNT(*) n, MAX(ts) last FROM punches WHERE method='device' GROUP BY source")}
    devices = []
    for d in list_devices():
        n, last = stats.get(d["dev_id"], (0, None))
        devices.append({**d, "total_logs": n, "last_punch": last})
    gps = [dict(r) for r in c.execute(
        "SELECT p.id, p.ts, p.device_user_id, COALESCE(e.name,'Unknown') name, p.selfie, p.lat, p.lng, p.flag, s.name site "
        "FROM punches p LEFT JOIN employees e ON e.device_user_id=p.device_user_id LEFT JOIN sites s ON s.id=p.site_id "
        "WHERE p.method='app' AND date(p.ts)=? AND COALESCE(p.flag,'')!='rejected'" + access.scope_sql("e.department")[0] +
        " ORDER BY p.ts DESC LIMIT 10", (day, *access.scope_sql("e.department")[1]))]
    recent = [dict(r) for r in c.execute(
        "SELECT p.ts, p.device_user_id, COALESCE(e.name,p.device_user_id) name, p.method, p.source, p.flag "
        "FROM punches p LEFT JOIN employees e ON e.device_user_id=p.device_user_id "
        "WHERE date(p.ts)=? AND COALESCE(p.flag,'')!='rejected'" + access.scope_sql("e.department")[0] +
        " ORDER BY p.ts DESC LIMIT 8", (day, *access.scope_sql("e.department")[1]))]
    pending = c.execute(
        "SELECT COUNT(*) FROM leaves l LEFT JOIN employees e ON e.device_user_id=l.device_user_id WHERE l.status='pending'"
        + access.scope_sql("e.department")[0], access.scope_sql("e.department")[1]).fetchone()[0]
    pending_wfh = c.execute(
        "SELECT COUNT(*) FROM wfh_requests w LEFT JOIN employees e ON e.device_user_id=w.device_user_id WHERE w.status='pending'"
        + access.scope_sql("e.department")[0], access.scope_sql("e.department")[1]).fetchone()[0]
    return {"day": day, "total": len(rows), "pending_wfh": pending_wfh,
            "on_leave": sum(r["status"] == "Leave" for r in rows), "recent": recent, "pending_leaves": pending,
            "present": sum(r["status"] in reports.WORKED for r in rows),
            "absent": sum(r["status"] == "Absent" for r in rows),
            "late": sum(r["status"] == "Late" for r in rows),
            "informed_leave": reports.leave_days(rows, "Informed"), "adhoc_leave": reports.leave_days(rows, "Adhoc"),
            "early_in": sum(r["early_in_min"] > 0 for r in rows), "early_out": sum(r["early_min"] > 0 for r in rows),
            "total_hours": round(sum(r["hours"] for r in rows), 2), "department": department, "dept_options": dept_options,
            "late_policy": lates.policy(), "late_total": sum(x["count"] for x in watch),
            "late_watch": [{k: x[k] for k in ("device_user_id", "name", "department", "count", "limit", "state")}
                           for x in watch if x["state"] != "ok"],
            "departments": [{"name": k, "present": v} for k, v in sorted(depts.items())],
            "devices": devices, "gps_punches": gps}


def _valid_date(d: str):
    try:
        date.fromisoformat(d)
    except ValueError:
        raise HTTPException(400, f"Invalid date {d!r} (use YYYY-MM-DD)")


@app.get("/api/shifts", dependencies=[Depends(access.need("manager"))])
def list_shifts():
    return [dict(r) for r in db.conn().execute(
        "SELECT s.*, (SELECT COUNT(*) FROM employees e WHERE e.shift_id=s.id AND e.active=1) employees FROM shifts s ORDER BY s.name")]


@app.post("/api/shifts")
async def save_shift(sh: ShiftIn, u: dict = Depends(access.need("manager", "shifts"))):
    if sh.id and u["role"] == "manager" and not access.has_perm(u, "shifts"):
        raise HTTPException(403, "Team leads can create new shifts but not change existing ones")
    if not (re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", sh.start) and re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", sh.end)):
        raise HTTPException(400, "Times must be HH:MM (24-hour)")
    if sh.end <= sh.start:
        raise HTTPException(400, "Shift end must be after start (overnight shifts are not supported yet)")
    offs = [o.strip().upper() for o in sh.weekly_off.split(",") if o.strip()]
    if any(o not in reports.WEEKDAYS for o in offs):
        raise HTTPException(400, "Weekly off must be a comma list of MON,TUE,WED,THU,FRI,SAT,SUN")
    if sh.break_start or sh.break_end:
        if not all(re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", t) for t in (sh.break_start, sh.break_end)):
            raise HTTPException(400, "Break times must be HH:MM (24-hour), or leave both empty")
        if not (sh.start <= sh.break_start < sh.break_end <= sh.end):
            raise HTTPException(400, "The break must be inside the shift and end after it starts")
    c = db.conn()
    vals = (sh.name.strip(), sh.start, sh.end, sh.grace_min, sh.half_day_hours, sh.ot_after_min, ",".join(offs),
            sh.break_start, sh.break_end)
    if sh.id:
        c.execute("UPDATE shifts SET name=?, start=?, end=?, grace_min=?, half_day_hours=?, ot_after_min=?, weekly_off=?, "
                  "break_start=?, break_end=? WHERE id=?", vals + (sh.id,))
    else:
        c.execute("INSERT INTO shifts(name,start,end,grace_min,half_day_hours,ot_after_min,weekly_off,break_start,break_end) "
                  "VALUES (?,?,?,?,?,?,?,?,?)", vals)
    c.commit()
    return {"ok": True}


@app.delete("/api/shifts/{shift_id}", dependencies=[Depends(access.need("hr", "shifts"))])
def delete_shift(shift_id: int):
    c = db.conn()
    c.execute("UPDATE employees SET shift_id=NULL WHERE shift_id=?", (shift_id,))
    c.execute("DELETE FROM shifts WHERE id=?", (shift_id,))
    c.commit()
    return {"ok": True}


@app.get("/api/shift-changes", dependencies=[Depends(access.need("manager"))])
def list_shift_changes(month: str | None = None):
    month = _valid_month(month or now_local().strftime("%Y-%m"))
    frag, args = access.scope_sql("e.department")
    return [dict(r) for r in db.conn().execute(
        "SELECT c.*, COALESCE(e.name, c.device_user_id) name, s.name shift, s.start, s.end FROM shift_changes c "
        "JOIN employees e ON e.device_user_id=c.device_user_id LEFT JOIN shifts s ON s.id=c.shift_id "
        "WHERE substr(c.start_date,1,7)<=? AND substr(c.end_date,1,7)>=?" + frag + " ORDER BY c.start_date DESC, c.id DESC",
        (month, month, *args))]


@app.post("/api/shift-changes")
async def add_shift_change(body: ShiftChangeIn, u: dict = Depends(access.need("manager", "shifts"))):
    """Put people on a different shift for a date range (a team lead can do this for their own team)."""
    _valid_date(body.start_date)
    _valid_date(body.end_date)
    if body.end_date < body.start_date:
        raise HTTPException(400, "End date is before start date")
    c = db.conn()
    if not c.execute("SELECT 1 FROM shifts WHERE id=?", (body.shift_id,)).fetchone():
        raise HTTPException(404, "No such shift")
    ids = [db.canon_id(i) for i in body.ids]
    if not ids:
        raise HTTPException(400, "Select at least one employee")
    for i in ids:
        if not c.execute("SELECT 1 FROM employees WHERE device_user_id=? AND active=1", (i,)).fetchone():
            raise HTTPException(404, f"No such employee: {i}")
        access.assert_in_scope(i)
    c.executemany("INSERT INTO shift_changes(device_user_id,start_date,end_date,shift_id,created_by) VALUES (?,?,?,?,?)",
                  [(i, body.start_date, body.end_date, body.shift_id, u["username"]) for i in ids])
    c.commit()
    return {"created": len(ids)}


@app.delete("/api/shift-changes/{change_id}", dependencies=[Depends(access.need("manager", "shifts"))])
def delete_shift_change(change_id: int):
    c = db.conn()
    row = c.execute("SELECT device_user_id FROM shift_changes WHERE id=?", (change_id,)).fetchone()
    if not row:
        raise HTTPException(404, "No such shift change")
    access.assert_in_scope(row["device_user_id"])
    c.execute("DELETE FROM shift_changes WHERE id=?", (change_id,))
    c.commit()
    return {"ok": True}


@app.get("/api/holidays", dependencies=[Depends(access.need("manager"))])
def list_holidays(year: str | None = None):
    year = year or now_local().strftime("%Y")
    return [dict(r) for r in db.conn().execute("SELECT * FROM holidays WHERE substr(date,1,4)=? ORDER BY date", (year,))]


@app.post("/api/holidays", dependencies=[Depends(access.need("hr", "holidays"))])
def save_holiday(h: HolidayIn):
    _valid_date(h.date)
    c = db.conn()
    c.execute("INSERT INTO holidays(date,name,optional) VALUES (?,?,?) ON CONFLICT(date) DO UPDATE SET name=excluded.name, "
              "optional=excluded.optional", (h.date, h.name, 0))
    c.commit()
    return {"ok": True}


@app.delete("/api/holidays/{day}", dependencies=[Depends(access.need("hr", "holidays"))])
def delete_holiday(day: str):
    c = db.conn()
    c.execute("DELETE FROM holidays WHERE date=?", (day,))
    c.commit()
    return {"ok": True}


@app.get("/api/leaves", dependencies=[Depends(access.need("manager"))])
def list_leaves(month: str | None = None):
    month = month or now_local().strftime("%Y-%m")
    frag, args = access.scope_sql("e.department")
    return [dict(r) for r in db.conn().execute(
        "SELECT l.*, COALESCE(e.name, l.device_user_id) name, COALESCE(e.department,'') department, "
        "(SELECT id FROM leave_consents k WHERE k.group_id=l.group_id) consent_id FROM leaves l LEFT JOIN employees e ON e.device_user_id=l.device_user_id "
        "WHERE substr(l.start_date,1,7)<=? AND substr(l.end_date,1,7)>=?" + frag + " ORDER BY l.start_date DESC", (month, month, *args))]


@app.post("/api/leaves", dependencies=[Depends(access.need("hr", "leave"))])
def add_leave(l: LeaveIn):
    l.device_user_id = db.canon_id(l.device_user_id)
    access.assert_in_scope(l.device_user_id)
    _valid_date(l.start_date)
    _valid_date(l.end_date)
    if l.end_date < l.start_date:
        raise HTTPException(400, "End date is before start date")
    l.type = leavepolicy.ALIASES.get(l.type, l.type)
    if l.type not in reports.LEAVE_TYPES:
        raise HTTPException(400, "Leave type must be one of " + ", ".join(reports.LEAVE_TYPES))
    c = db.conn()
    if not c.execute("SELECT 1 FROM employees WHERE device_user_id=?", (l.device_user_id,)).fetchone():
        raise HTTPException(404, "No such employee")
    today = now_local().date()
    if l.applied_on:
        _valid_date(l.applied_on)
    applied = date.fromisoformat(l.applied_on) if l.applied_on else today
    if l.half and l.end_date != l.start_date:
        raise HTTPException(400, "A half day is a single date")
    parts = [{"type": l.type, "lop": 0, "half": int(l.half), "start_date": l.start_date, "end_date": l.end_date}]
    if l.type in leavepolicy.POLICY_TYPES and not l.override:   # split by the leave policy, like the phone app
        p = leavepolicy.plan(l.device_user_id, l.type, l.start_date, l.end_date, applied, strict=False, half=l.half)
        if "error" in p:
            raise HTTPException(400, p["error"] + " (tick 'override' to record it anyway)")
        parts = p["parts"]
    group = secrets.token_hex(6)
    for part in parts:
        c.execute("INSERT INTO leaves(device_user_id,start_date,end_date,type,note,status,applied_on,requested_by,group_id,lop,half) "
                  "VALUES (?,?,?,?,?,?,?,?,?,?,?)", (l.device_user_id, part["start_date"], part["end_date"], part["type"], l.note,
                                                    "approved", applied.isoformat(), "", group, part["lop"], part["half"]))
    c.commit()
    return {"ok": True, "parts": parts}


class AbsenceRuleIn(BaseModel):
    start: str


@app.get("/api/absence-rule", dependencies=[Depends(access.need("manager"))])
def get_absence_rule():
    return {"start": leavepolicy.absence_rule_from().isoformat()}


@app.post("/api/absence-rule", dependencies=[Depends(access.need("hr", "leave"))])
def set_absence_rule(body: AbsenceRuleIn):
    """From which day an absence without approved leave counts as Adhoc (e.g. backdate it to the 1st of the month)."""
    _valid_date(body.start)
    c = db.conn()
    c.execute("INSERT INTO settings(key,value) VALUES ('absent_as_adhoc_from',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
              (body.start,))
    c.commit()
    return {"start": body.start}


class LeaveAdjustIn(BaseModel):
    device_user_id: str
    fy: int
    informed: float = 0
    adhoc: float = 0
    reason: str


@app.get("/api/leave-adjustments", dependencies=[Depends(access.need("hr", "leave"))])
def list_leave_adjustments(fy: int):
    frag, args = access.scope_sql("e.department")
    return [dict(r) for r in db.conn().execute(
        "SELECT a.*, COALESCE(e.name, a.device_user_id) name, COALESCE(e.department,'') department FROM leave_adjustments a "
        "LEFT JOIN employees e ON e.device_user_id=a.device_user_id WHERE a.fy=?" + frag + " ORDER BY a.id DESC", [fy] + args)]


@app.post("/api/leave-adjustments")
async def add_leave_adjustment(body: LeaveAdjustIn, u: dict = Depends(access.need("hr", "leave"))):
    """Give someone extra leave days for a financial year (or take some away), e.g. a balance carried over from the old
    system. They count in the balance and in the year-end encashment."""
    uid = db.canon_id(body.device_user_id)
    access.assert_in_scope(uid)
    c = db.conn()
    if not c.execute("SELECT 1 FROM employees WHERE device_user_id=?", (uid,)).fetchone():
        raise HTTPException(404, "No such employee")
    if not 2000 <= body.fy <= 2100:
        raise HTTPException(400, "Choose the financial year")
    for v in (body.informed, body.adhoc):
        if abs(v) > 60 or (v * 2) != int(v * 2):
            raise HTTPException(400, "Days must be whole or half days, at most 60")
    if not (body.informed or body.adhoc):
        raise HTTPException(400, "Enter the Informed and / or Adhoc days to add (a minus number takes days away)")
    if not body.reason.strip():
        raise HTTPException(400, "Give a reason (it is kept with the adjustment)")
    cur = c.execute("INSERT INTO leave_adjustments(device_user_id,fy,informed,adhoc,reason,created_by,created_at) VALUES (?,?,?,?,?,?,?)",
                    (uid, body.fy, body.informed, body.adhoc, body.reason.strip()[:200], u["username"],
                     now_local().strftime("%Y-%m-%d %H:%M:%S")))
    c.commit()
    return {"ok": True, "id": cur.lastrowid}


@app.delete("/api/leave-adjustments/{adj_id}", dependencies=[Depends(access.need("hr", "leave"))])
def delete_leave_adjustment(adj_id: int):
    c = db.conn()
    row = c.execute("SELECT device_user_id FROM leave_adjustments WHERE id=?", (adj_id,)).fetchone()
    if not row:
        raise HTTPException(404, "No such adjustment")
    access.assert_in_scope(row["device_user_id"])
    c.execute("DELETE FROM leave_adjustments WHERE id=?", (adj_id,))
    c.commit()
    return {"ok": True}


@app.get("/api/leave-preview", dependencies=[Depends(access.need("hr", "leave"))])
def leave_preview(employee: str, type: str, start_date: str, end_date: str = "", applied_on: str = "", half: bool = False):
    uid = db.canon_id(employee)
    access.assert_in_scope(uid)
    applied = date.fromisoformat(applied_on) if applied_on else now_local().date()
    p = leavepolicy.plan(uid, type, start_date, end_date or start_date, applied, strict=False, half=half)
    if "error" in p:
        raise HTTPException(400, p["error"])
    return p


@app.get("/api/wfh", dependencies=[Depends(access.need("manager"))])
def list_wfh(month: str | None = None, status: str | None = None):
    month = _valid_month(month or now_local().strftime("%Y-%m"))
    frag, args = access.scope_sql("e.department")
    q = ("SELECT w.*, COALESCE(e.name, w.device_user_id) name, COALESCE(e.department,'') department FROM wfh_requests w "
         "LEFT JOIN employees e ON e.device_user_id=w.device_user_id WHERE substr(w.start_date,1,7)<=? AND substr(w.end_date,1,7)>=?" + frag)
    args = [month, month, *args]
    if status:
        q, args = q + " AND w.status=?", args + [status]
    return [dict(r) for r in db.conn().execute(q + " ORDER BY w.status='pending' DESC, w.start_date DESC", args)]


@app.post("/api/wfh/{req_id}/{action}")
async def decide_wfh(req_id: int, action: str, u: dict = Depends(access.need("manager", "leave"))):
    """Team leads (own team), HR, management and admin approve or reject work-from-home requests."""
    if action not in ("approve", "reject"):
        raise HTTPException(404)
    c = db.conn()
    row = c.execute("SELECT * FROM wfh_requests WHERE id=?", (req_id,)).fetchone()
    if not row:
        raise HTTPException(404, "No such request")
    access.assert_in_scope(row["device_user_id"])
    if row["status"] != "pending":
        raise HTTPException(400, f"This request is already {row['status']}")
    c.execute("UPDATE wfh_requests SET status=?, decided_by=?, decided_at=? WHERE id=?",
              ("approved" if action == "approve" else "rejected", u["username"], now_local().strftime("%Y-%m-%d %H:%M:%S"), req_id))
    c.commit()
    return {"ok": True}


@app.get("/api/leave-summary", dependencies=[Depends(access.need("hr", "leave"))])
def leave_summary(fy: int | None = None):
    """Leave used this financial year and the balance to encash on 31 March."""
    return leavepolicy.year_summary(fy if fy else leavepolicy.fy(now_local().date()))


@app.get("/api/leave-consents/{consent_id}", dependencies=[Depends(access.need("hr", "leave"))])
def leave_consent(consent_id: int):
    row = db.conn().execute("SELECT k.*, COALESCE(e.name,k.device_user_id) name FROM leave_consents k "
                            "LEFT JOIN employees e ON e.device_user_id=k.device_user_id WHERE k.id=?", (consent_id,)).fetchone()
    if not row:
        raise HTTPException(404, "No such consent")
    access.assert_in_scope(row["device_user_id"])
    return dict(row)


@app.get("/api/leave-consents/{consent_id}/pdf", dependencies=[Depends(access.need("hr", "leave"))])
def leave_consent_pdf(consent_id: int):
    k = leave_consent(consent_id)
    parts = [dict(r) for r in db.conn().execute("SELECT * FROM leaves WHERE group_id=? ORDER BY start_date", (k["group_id"],))]
    return Response(payslip.consent_pdf(k, parts, config.COMPANY_NAME), media_type="application/pdf",
                    headers={"Content-Disposition": f"inline; filename=leave-consent-{consent_id}.pdf"})


@app.post("/api/leaves/{leave_id}/approve")
async def approve_leave(leave_id: int, u: dict = Depends(access.need("hr", "leave"))):
    return _decide_leave(leave_id, "approved", u)


@app.post("/api/leaves/{leave_id}/reject")
async def reject_leave(leave_id: int, u: dict = Depends(access.need("hr", "leave"))):
    return _decide_leave(leave_id, "rejected", u)


def _decide_leave(leave_id: int, status: str, u: dict) -> dict:
    c = db.conn()
    row = c.execute("SELECT * FROM leaves WHERE id=?", (leave_id,)).fetchone()
    if not row:
        raise HTTPException(404, "No such leave")
    access.assert_in_scope(row["device_user_id"])
    if row["status"] != "pending":
        raise HTTPException(400, f"This leave is already {row['status']}")
    # an application split into Informed / Adhoc parts is decided as a whole
    c.execute("UPDATE leaves SET status=?, decided_by=?, decided_at=? WHERE status='pending' AND (id=? OR group_id=?)",
              (status, u["username"], now_local().strftime("%Y-%m-%d %H:%M:%S"), leave_id, row["group_id"] or "-"))
    c.commit()
    return {"ok": True}


@app.get("/api/leave-balance", dependencies=[Depends(access.need("manager"))])
def leave_balance(employee: str, on: str | None = None):
    uid = db.canon_id(employee)
    access.assert_in_scope(uid)
    if on:
        _valid_date(on)
    b = leavepolicy.balance(uid, date.fromisoformat(on) if on else now_local().date())
    if b is None:
        raise HTTPException(404, "No such employee")
    return b


@app.delete("/api/leaves/{leave_id}", dependencies=[Depends(access.need("hr", "leave"))])
def delete_leave(leave_id: int):
    c = db.conn()
    row = c.execute("SELECT device_user_id, group_id FROM leaves WHERE id=?", (leave_id,)).fetchone()
    if row:
        access.assert_in_scope(row["device_user_id"])
        c.execute("DELETE FROM leaves WHERE id=? OR group_id=?", (leave_id, row["group_id"] or "-"))
    c.commit()
    return {"ok": True}


MAX_IMPORT_BYTES = 15 * 1024 * 1024


@app.post("/api/import/{kind}", dependencies=[Depends(access.need("hr", "import"))])
async def import_file(kind: str, file: UploadFile = File(...), dry_run: bool = True, inactive: bool = False):
    data = await file.read(MAX_IMPORT_BYTES + 1)
    if len(data) > MAX_IMPORT_BYTES:
        raise HTTPException(413, "File too large (max 15 MB)")
    try:
        return importer.run(kind, file.filename or "", data, dry_run, inactive)
    except importer.ImportError_ as e:
        raise HTTPException(400, str(e))
    except Exception as e:  # corrupt xlsx etc.: report, never a stack trace
        logging.exception("import failed")
        raise HTTPException(400, f"Could not read the file: {type(e).__name__}")


@app.get("/api/import/{kind}/template.csv", dependencies=[Depends(access.need("hr", "import"))])
def import_template(kind: str):
    if kind not in importer.TEMPLATES:
        raise HTTPException(404)
    return PlainTextResponse(importer.template_csv(kind), media_type="text/csv",
                             headers={"Content-Disposition": f"attachment; filename={kind}-template.csv"})


@app.get("/api/sites", dependencies=[Depends(access.need("manager"))])
def list_sites():
    return [dict(r) for r in db.conn().execute(
        "SELECT s.*, COALESCE(e.name,'') employee FROM sites s LEFT JOIN employees e ON e.device_user_id=s.device_user_id "
        "WHERE s.active=1 ORDER BY s.kind DESC, s.name")]


@app.post("/api/sites", dependencies=[Depends(access.need("admin", "sites"))])
def add_site(s: SiteIn):
    if not (-90 <= s.lat <= 90 and -180 <= s.lng <= 180 and s.radius_m > 0):
        raise HTTPException(400, "Invalid latitude/longitude/radius")
    c = db.conn()
    uid = None
    if s.kind == "home":
        uid = db.canon_id(s.device_user_id)
        if not c.execute("SELECT 1 FROM employees WHERE device_user_id=?", (uid,)).fetchone():
            raise HTTPException(400, "Choose the employee whose home this is")
    elif s.kind != "office":
        raise HTTPException(400, "Site type must be office or home")
    c.execute("INSERT INTO sites(name,lat,lng,radius_m,kind,device_user_id) VALUES (?,?,?,?,?,?)",
              (s.name, s.lat, s.lng, s.radius_m, s.kind, uid))
    c.commit()
    return {"ok": True}


@app.delete("/api/sites/{site_id}", dependencies=[Depends(access.need("admin", "sites"))])
def delete_site(site_id: int):
    c = db.conn()
    c.execute("UPDATE sites SET active=0 WHERE id=?", (site_id,))
    c.commit()
    return {"ok": True}


@app.get("/api/punches", dependencies=[Depends(access.need("manager"))])
def list_punches(day: str | None = None, limit: int = 500, device: str | None = None):
    day = day or now_local().date().isoformat()
    frag, args = access.scope_sql("e.department")
    rows = db.conn().execute(
        "SELECT p.id, p.ts, p.device_user_id, COALESCE(e.name,'Unknown') name, p.status, p.source, p.method, "
        "p.flag, p.note, p.selfie, p.lat, p.lng, p.accuracy, p.distance_m, s.name site FROM punches p "
        "LEFT JOIN employees e ON e.device_user_id=p.device_user_id LEFT JOIN sites s ON s.id=p.site_id "
        "WHERE date(p.ts)=? AND (? IS NULL OR p.source=?)" + frag + " ORDER BY p.ts DESC LIMIT ?",
        (day, device, device, *args, limit),
    )
    return [dict(r) for r in rows]


@app.post("/api/punches/{punch_id}/review", dependencies=[Depends(access.need("manager", "punches"))])
def review_punch(punch_id: int, body: ReviewIn):
    if body.flag not in ("rejected", "approved", ""):
        raise HTTPException(400, "flag must be rejected, approved or empty")
    c = db.conn()
    row = c.execute("SELECT device_user_id FROM punches WHERE id=?", (punch_id,)).fetchone()
    if not row:
        raise HTTPException(404, "No such punch")
    access.assert_in_scope(row["device_user_id"])
    c.execute("UPDATE punches SET flag=? WHERE id=?", (body.flag, punch_id))
    c.commit()
    return {"ok": True}


@app.get("/api/selfies/{year_month}/{filename}", dependencies=[Depends(access.need("manager"))])
def get_selfie(year_month: str, filename: str):
    owner = db.conn().execute("SELECT device_user_id FROM punches WHERE selfie=?", (f"{year_month}/{filename}",)).fetchone()
    if owner:
        access.assert_in_scope(owner["device_user_id"])
    base = os.path.realpath(config.SELFIE_DIR)
    path = os.path.realpath(os.path.join(base, year_month, filename))
    if not path.startswith(base + os.sep) or not os.path.isfile(path):
        raise HTTPException(404)
    return FileResponse(path, headers={"Cache-Control": "private, max-age=86400"})


@app.post("/api/punches/day", dependencies=[Depends(access.need("manager", "punches"))])
def manual_day(p: DayPunchIn):
    """Enter or fix a day's punch-in and/or punch-out (HH:MM, 24-hour) for an employee."""
    uid = db.canon_id(p.employee_id)
    _valid_date(p.date)
    if p.date > now_local().date().isoformat():
        raise HTTPException(400, "You cannot enter punches for a future date")
    times = [t.strip() for t in (p.in_time, p.out_time) if t.strip()]
    if not times:
        raise HTTPException(400, "Enter a punch-in time, a punch-out time, or both")
    if not db.conn().execute("SELECT 1 FROM employees WHERE device_user_id=?", (uid,)).fetchone():
        raise HTTPException(404, "No such employee")
    access.assert_in_scope(uid)
    stored = 0
    for kind, t in (("in", p.in_time.strip()), ("out", p.out_time.strip())):
        if not t:
            continue
        if not re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", t):
            raise HTTPException(400, "Times must be HH:MM in 24-hour format, e.g. 09:30")
        rec = db.add_punch(uid, f"{p.date} {t}:00", 0 if kind == "in" else 1, "manual", method="manual", note=(p.remark or "")[:45])
        if rec:
            stored += 1
            hub.publish(rec)
    if not stored:
        raise HTTPException(409, "Nothing saved: a punch already exists at (or within a minute of) that time")
    return {"stored": stored}


@app.post("/api/punches/{punch_id}/correct", dependencies=[Depends(access.need("manager", "punches"))])
def correct_punch(punch_id: int, body: CorrectIn):
    """Replace a punch's time: the old punch is kept but rejected, and a manual punch is added at the new time."""
    c = db.conn()
    row = c.execute("SELECT * FROM punches WHERE id=?", (punch_id,)).fetchone()
    if not row:
        raise HTTPException(404, "No such punch")
    access.assert_in_scope(row["device_user_id"])
    ts = body.ts.strip().replace("T", " ")
    ts = ts + ":00" if re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}", ts) else ts
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2} ([01]\d|2[0-3]):[0-5]\d:[0-5]\d", ts):
        raise HTTPException(400, "Time must look like 2026-10-05 09:30")
    _valid_date(ts[:10])
    c.execute("UPDATE punches SET flag='rejected' WHERE id=?", (punch_id,))
    c.commit()
    rec = db.add_punch(row["device_user_id"], ts, row["status"], "manual", method="manual",
                       note=(f"corrected from {row['ts'][11:16]}" + (f": {body.remark}" if body.remark else ""))[:45])
    if rec is None:   # same time already exists: undo, nothing changes
        c.execute("UPDATE punches SET flag=? WHERE id=?", (row["flag"] or "", punch_id))
        c.commit()
        raise HTTPException(409, "A punch already exists at that time")
    hub.publish(rec)
    return {"ok": True}


@app.post("/api/punches", dependencies=[Depends(access.need("manager", "punches"))])
def manual_punch(p: PunchIn):
    access.assert_in_scope(db.canon_id(p.device_user_id))
    rec = db.add_punch(p.device_user_id, p.ts, p.status, "manual", method="manual")
    if rec:
        hub.publish(rec)
    return {"stored": rec is not None}


def _csv(rows: list[dict], filename: str) -> PlainTextResponse:
    buf = io.StringIO()
    if rows:
        w = csv.DictWriter(buf, fieldnames=rows[0].keys())
        w.writeheader()
        w.writerows([{k: _safe(v) for k, v in r.items()} for r in rows])
    return PlainTextResponse(buf.getvalue(), media_type="text/csv",
                             headers={"Content-Disposition": f"attachment; filename={filename}"})


def _xlsx(rows: list[dict], filename: str, title: str = "Report") -> Response:
    from openpyxl import Workbook
    wb = Workbook()
    xlsx_sheet(wb.active, title, rows)
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type=XLSX_TYPE, headers={"Content-Disposition": f"attachment; filename={filename}"})


def _grid_rows(g: dict) -> list[dict]:
    return [{"ID": r["device_user_id"], "Name": r["name"], "Department": r["department"],
             **{str(d): c for d, c in zip(g["days"], r["codes"])}, **{k: v for k, v in r["totals"].items()}} for r in g["rows"]]


def _day(day): return day or now_local().date().isoformat()
def _month(month): return month or now_local().strftime("%Y-%m")


@app.get("/api/reports/daily", dependencies=[Depends(access.need("manager"))])
def report_daily(day: str | None = None, device: str | None = None):
    return reports.daily(_day(day), device)


@app.get("/api/reports/daily.csv", dependencies=[Depends(access.need("manager"))])
def report_daily_csv(day: str | None = None, device: str | None = None):
    return _csv(reports.daily(_day(day), device), "attendance-daily.csv")


@app.get("/api/reports/monthly", dependencies=[Depends(access.need("manager"))])
def report_monthly(month: str | None = None, device: str | None = None):
    return reports.monthly(_month(month), device)


@app.get("/api/reports/monthly.csv", dependencies=[Depends(access.need("manager"))])
def report_monthly_csv(month: str | None = None, device: str | None = None):
    return _csv(reports.monthly(_month(month), device), "attendance-monthly.csv")


@app.get("/api/reports/grid", dependencies=[Depends(access.need("manager"))])
def report_grid(month: str | None = None, device: str | None = None):
    return reports.grid(_month(month), device)


@app.get("/api/reports/grid.csv", dependencies=[Depends(access.need("manager"))])
def report_grid_csv(month: str | None = None, device: str | None = None):
    return _csv(_grid_rows(reports.grid(_month(month), device)), "attendance-sheet.csv")


@app.get("/api/reports/daily.xlsx", dependencies=[Depends(access.need("manager"))])
def report_daily_xlsx(day: str | None = None, device: str | None = None):
    return _xlsx(reports.daily(_day(day), device), "attendance-daily.xlsx", "Daily")


@app.get("/api/reports/monthly.xlsx", dependencies=[Depends(access.need("manager"))])
def report_monthly_xlsx(month: str | None = None, device: str | None = None):
    return _xlsx(reports.monthly(_month(month), device), "attendance-monthly.xlsx", "Monthly")


@app.get("/api/reports/grid.xlsx", dependencies=[Depends(access.need("manager"))])
def report_grid_xlsx(month: str | None = None, device: str | None = None):
    return _xlsx(_grid_rows(reports.grid(_month(month), device)), "attendance-sheet.xlsx", "Sheet")


@app.get("/api/reports/payroll.xlsx", dependencies=[Depends(access.need("hr", "payroll"))])
def report_payroll_xlsx(month: str | None = None, device: str | None = None):
    return _xlsx(reports.payroll(_month(month), device), "payroll.xlsx", "Payroll")


@app.get("/api/reports/muster", dependencies=[Depends(access.need("manager"))])
def report_muster(month: str | None = None, device: str | None = None):
    return reports.muster(_month(month), device)


@app.get("/api/reports/muster.csv", dependencies=[Depends(access.need("manager"))])
def report_muster_csv(month: str | None = None, device: str | None = None):
    return _csv(reports.muster(_month(month), device), "muster-roll.csv")


@app.get("/api/reports/muster.xlsx", dependencies=[Depends(access.need("manager"))])
def report_muster_xlsx(month: str | None = None, device: str | None = None):
    return _xlsx(reports.muster(_month(month), device), "muster-roll.xlsx", "Muster Roll")


@app.get("/api/reports/timecard", dependencies=[Depends(access.need("manager"))])
def report_timecard(employee: str, month: str | None = None):
    return reports.timecard(_month(month), db.canon_id(employee))


def _range(start: str, end: str) -> tuple[str, str]:
    for d in (start, end):
        _valid_date(d)
    if end < start:
        raise HTTPException(400, "The end date is before the start date")
    if (date.fromisoformat(end) - date.fromisoformat(start)).days >= reports.MAX_TIMECARD_DAYS:
        raise HTTPException(400, "Choose at most one year")
    return start, end


def _timecards(start: str, end: str, employee: str, department: str, status: str) -> list[dict]:
    if status and status not in reports.TIMECARD_STATUSES:
        raise HTTPException(400, "Unknown status filter")
    uid = db.canon_id(employee) if employee else None
    if uid:
        access.assert_in_scope(uid)
    return reports.timecards(*_range(start, end), uid, department or None, status or None)


class LatePolicyIn(BaseModel):
    period: str = "month"
    limit: int = 3


@app.get("/api/late-policy", dependencies=[Depends(access.need("manager"))])
def get_late_policy():
    return {**lates.policy(), "periods": lates.PERIODS, "emails": config.LATE_EMAILS and bool(config.SMTP_HOST),
            "hr_email": config.HR_EMAIL}


@app.post("/api/late-policy", dependencies=[Depends(access.need("hr", "shifts"))])
def set_late_policy(body: LatePolicyIn):
    try:
        return lates.save_policy(body.period, body.limit)
    except ValueError as e:
        raise HTTPException(400, str(e))


@app.get("/api/lates", dependencies=[Depends(access.need("manager"))])
def late_counts(as_of: str = ""):
    """Qualifying late arrivals per employee in the current tracking period (or the one containing as_of)."""
    if as_of:
        _valid_date(as_of)
    return lates.counts(date.fromisoformat(as_of) if as_of else None)


@app.get("/api/late-notices", dependencies=[Depends(access.need("hr"))])
def late_notices(month: str = ""):
    month = _month(month)
    return [dict(r) for r in db.conn().execute(
        "SELECT n.*, COALESCE(e.name, n.device_user_id) name FROM late_notices n LEFT JOIN employees e "
        "ON e.device_user_id=n.device_user_id WHERE n.date LIKE ? ORDER BY n.date DESC, name", (month + "%",))]


@app.get("/api/reports/timecards", dependencies=[Depends(access.need("manager"))])
def report_timecards(start: str, end: str, employee: str = "", department: str = "", status: str = ""):
    """Time card for everyone (employee empty) or one person over a date range, optionally by department / status."""
    return _timecards(start, end, employee, department, status)


def _sheet_names(people: list[tuple[str, str]]) -> list[str]:
    """Excel sheet names: at most 31 characters, none of []:*?/\\, unique even for people with the same name."""
    used, out = {"all employees", "in-out"}, []
    for name, uid in people:
        clean = lambda t: re.sub(r"[\[\]:*?/\\]", "", t).strip("' ")
        tag, nm = f" ({clean(uid)})", clean(name) or "Employee"
        cand, n = nm[:31 - len(tag)].rstrip() + tag, 2
        while cand.lower() in used:   # same name and ID (should not happen) still gets its own sheet
            tail = f"{tag} ~{n}"
            cand, n = nm[:31 - len(tail)].rstrip() + tail, n + 1
        used.add(cand.lower())
        out.append(cand)
    return out


@app.get("/api/reports/timecards.xlsx", dependencies=[Depends(access.need("manager"))])
def report_timecards_xlsx(start: str, end: str, employee: str = "", department: str = "", status: str = ""):
    """One workbook: 'In-Out' (one row per employee, In and Out for each date), 'All Employees' (every row with the
    details, by date then name) and one sheet per employee with only their rows."""
    from openpyxl import Workbook
    rows = _timecards(start, end, employee, department, status)
    human = timecard_human
    wb = Workbook()
    inout_sheet(wb.active, "In-Out", rows, start, end)
    xlsx_sheet(wb.create_sheet(), "All Employees", [human(r) for r in rows] or [{label: "" for _, label in TIMECARD_COLUMNS}])
    people = sorted({(r["name"], r["device_user_id"]) for r in rows}, key=lambda p: (p[0].lower(), p[1]))
    for (name, uid), title in zip(people, _sheet_names(people)):
        mine = [human(r) for r in rows if r["device_user_id"] == uid]
        dept = mine[0]["Department"] if mine else ""
        xlsx_sheet(wb.create_sheet(), title, mine,
                   header_rows=[["Employee", name], ["Employee ID", uid], ["Department", dept], ["Period", f"{start} to {end}"]])
    out = io.BytesIO()
    wb.save(out)
    who = f"-{db.canon_id(employee)}" if employee else "-all"
    return Response(out.getvalue(), media_type=XLSX_TYPE,
                    headers={"Content-Disposition": f"attachment; filename=time-card{who}-{start}-to-{end}.xlsx"})


@app.get("/api/reports/timecard.xlsx", dependencies=[Depends(access.need("manager"))])
def report_timecard_xlsx(employee: str, month: str | None = None):
    return _xlsx(reports.timecard(_month(month), db.canon_id(employee)), "time-card.xlsx", "Time Card")


MONTH_RE = re.compile(r"\d{4}-(0[1-9]|1[0-2])")


def _valid_month(m: str) -> str:
    if not MONTH_RE.fullmatch(m or ""):
        raise HTTPException(400, "Month must look like 2026-09")
    return m


def _published(month: str) -> bool:
    row = db.conn().execute("SELECT published FROM payroll_months WHERE month=?", (month,)).fetchone()
    return bool(row and row["published"])


class AdjustmentIn(BaseModel):
    month: str = ""           # payroll month; taken from `date` when that is given
    date: str = ""            # the day the bonus / fine / ... was given
    device_user_id: str
    label: str
    amount: float  # positive = extra pay, negative = deduction


@app.get("/api/payroll/{month}/status", dependencies=[Depends(access.need("hr", "payroll"))])
def payroll_status(month: str):
    row = db.conn().execute("SELECT * FROM payroll_months WHERE month=?", (_valid_month(month),)).fetchone()
    return dict(row) if row else {"month": month, "published": 0, "published_by": "", "published_at": None}


@app.post("/api/payroll/{month}/publish")
async def payroll_publish(month: str, u: dict = Depends(access.need("hr", "payroll_edit"))):
    """Make payslips for a finished month visible to employees in the phone app."""
    _valid_month(month)
    if reports._month_days(month)[2] >= now_local().date().isoformat():
        raise HTTPException(400, "Publish only after the month has ended, so attendance is final")
    c = db.conn()
    c.execute("INSERT INTO payroll_months(month, published, published_by, published_at) VALUES (?,1,?,?) "
              "ON CONFLICT(month) DO UPDATE SET published=1, published_by=excluded.published_by, published_at=excluded.published_at",
              (month, u["username"], now_local().strftime("%Y-%m-%d %H:%M:%S")))
    c.commit()
    loans.record_month(month, u["username"])   # this month's EMIs become repayments in each loan's ledger
    return {"ok": True}


@app.post("/api/payroll/{month}/unpublish", dependencies=[Depends(access.need("hr", "payroll_edit"))])
def payroll_unpublish(month: str):
    c = db.conn()
    c.execute("UPDATE payroll_months SET published=0 WHERE month=?", (_valid_month(month),))
    c.commit()
    loans.unrecord_month(month)
    return {"ok": True}


@app.get("/api/payroll/adjustments", dependencies=[Depends(access.need("hr", "payroll"))])
def list_adjustments(month: str):
    return [dict(r) for r in db.conn().execute(
        "SELECT a.*, COALESCE(e.name, a.device_user_id) name FROM payroll_adjustments a "
        "LEFT JOIN employees e ON e.device_user_id=a.device_user_id WHERE a.month=? ORDER BY a.id", (_valid_month(month),))]


@app.post("/api/payroll/adjustments", dependencies=[Depends(access.need("hr", "payroll_edit"))])
def add_adjustment(a: AdjustmentIn):
    if a.date:
        _valid_date(a.date)
        a.month = a.date[:7]
    _valid_month(a.month)
    uid = db.canon_id(a.device_user_id)
    if not a.label.strip() or a.amount == 0 or abs(a.amount) > 10_000_000:
        raise HTTPException(400, "Give a label and a non-zero amount")
    if _published(a.month):
        raise HTTPException(400, "This month is published. Unpublish it before changing amounts.")
    c = db.conn()
    if not c.execute("SELECT 1 FROM employees WHERE device_user_id=?", (uid,)).fetchone():
        raise HTTPException(404, "No such employee")
    c.execute("INSERT INTO payroll_adjustments(month,device_user_id,label,amount,date) VALUES (?,?,?,?,?)",
              (a.month, uid, a.label.strip()[:60], a.amount, a.date or None))
    c.commit()
    return {"ok": True}


@app.delete("/api/payroll/adjustments/{adj_id}", dependencies=[Depends(access.need("hr", "payroll_edit"))])
def delete_adjustment(adj_id: int):
    c = db.conn()
    row = c.execute("SELECT month FROM payroll_adjustments WHERE id=?", (adj_id,)).fetchone()
    if row and _published(row["month"]):
        raise HTTPException(400, "This month is published. Unpublish it before changing amounts.")
    c.execute("DELETE FROM payroll_adjustments WHERE id=?", (adj_id,))
    c.commit()
    return {"ok": True}


@app.get("/api/payslips/{uid}/{month}.pdf", dependencies=[Depends(access.need("hr", "payroll"))])
def payslip_pdf(uid: str, month: str):
    uid = db.canon_id(uid)
    if not reports.payslip_ready(_valid_month(month), uid):
        raise HTTPException(400, "A payslip can be made only after the month has ended, or for someone who has "
                                 "left (set their last working day in Employees)")
    d = reports.payslip(month, uid)
    if not d:
        raise HTTPException(404, "No payroll for this employee and month")
    return Response(payslip.build_pdf(d), media_type="application/pdf",
                    headers={"Content-Disposition": f"inline; filename=payslip-{uid}-{month}.pdf"})


@app.get("/api/reports/payroll", dependencies=[Depends(access.need("hr", "payroll"))])
def report_payroll(month: str | None = None, device: str | None = None):
    return reports.payroll(_month(month), device)


@app.get("/api/reports/payroll.csv", dependencies=[Depends(access.need("hr", "payroll"))])
def report_payroll_csv(month: str | None = None, device: str | None = None):
    return _csv(reports.payroll(_month(month), device), "payroll.csv")


@app.post("/api/devices/sync", dependencies=[Depends(access.need("admin", "devices"))])
def sync_now():
    """Trigger an immediate pull from all configured ZK devices."""
    return {a: zk_poller.sync_device(a) for a in config.ZK_DEVICES}


@app.websocket("/ws")
async def ws(websocket: WebSocket):
    if not access.open_mode():   # live feed shows names: logged-in hr/admin only
        user = access.session_user(websocket)
        if not user or access.LEVEL[user["role"]] < access.LEVEL["hr"]:
            await websocket.close(code=4401)
            return
    await websocket.accept()
    await hub.register(websocket)
    try:
        while True:
            await websocket.receive_text()
    except WebSocketDisconnect:
        hub.unregister(websocket)


@app.get("/")
def index(request: Request):
    if not access.open_mode() and not access.session_user(request):
        return RedirectResponse("/login")
    return FileResponse("app/static/index.html")


@app.get("/login")
def login_page():
    return FileResponse("app/static/login.html")


@app.get("/m")
def mobile_redirect():
    return RedirectResponse("/m/")


@app.get("/healthz")
def healthz():
    return {"ok": True}


@app.get("/api/mail", dependencies=[Depends(access.need("admin"))])
def mail_status():
    """E-mail set-up and the weekly report's last run, for the Backups page."""
    a, b = weekly.period(now_local().date())
    return {"configured": mailer.configured(), "host": config.SMTP_HOST, "port": config.SMTP_PORT,
            "sender": config.SMTP_FROM or config.SMTP_USER, "hr_email": config.HR_EMAIL,
            "weekly": {"enabled": config.WEEKLY_REPORT, "hour": config.WEEKLY_REPORT_HOUR, "last": weekly.status(),
                       "next_from": a.isoformat(), "next_to": b.isoformat()}}


@app.post("/api/mail/test", dependencies=[Depends(access.need("admin"))])
def mail_test():
    """Send a short test e-mail to HR; the error (never the password) comes back if it fails."""
    if not mailer.configured():
        raise HTTPException(400, "E-mail is not set up: fill in SMTP_HOST, SMTP_PORT, SMTP_USER and SMTP_PASSWORD in .env and restart")
    try:
        mailer.send([config.HR_EMAIL], "Attendance system: test e-mail",
                    f"This is a test from the attendance system at {now_local():%d %b %Y %H:%M}.\n\n"
                    "If you can read this, e-mail is working: late-arrival alerts, the nightly backup and the weekly "
                    "report will reach you.")
    except Exception as e:
        raise HTTPException(400, f"Sending failed: {type(e).__name__}: {e}"[:400])
    return {"ok": True, "sent_to": config.HR_EMAIL}


@app.post("/api/weekly-report/run", dependencies=[Depends(access.need("admin"))])
def weekly_run():
    return weekly.run()


@app.get("/api/weekly-report.xlsx", dependencies=[Depends(access.need("admin"))])
def weekly_xlsx():
    a, b = weekly.period(now_local().date())
    data, _ = weekly.build(a, b)
    return Response(data, media_type=XLSX_TYPE,
                    headers={"Content-Disposition": f"attachment; filename=attendance-report-{a}-to-{b}.xlsx"})


@app.get("/api/backups", dependencies=[Depends(access.need("admin"))])
def list_backups():
    return {"last": backup.status(), "files": backup.list_files(), "email_to": config.BACKUP_EMAIL_TO,
            "email_configured": backup.email_configured(), "encrypted": bool(config.BACKUP_PASSWORD),
            "hour": config.BACKUP_HOUR, "keep_days": config.BACKUP_KEEP_DAYS}


@app.post("/api/backups/run", dependencies=[Depends(access.need("admin"))])
async def run_backup():
    return await asyncio.to_thread(backup.run, now_local().replace(microsecond=0), True)


@app.get("/api/backups/{name}", dependencies=[Depends(access.need("admin"))])
def download_backup(name: str):
    if not backup.NAME_RE.fullmatch(name) or not os.path.isfile(os.path.join(config.BACKUP_DIR, name)):
        raise HTTPException(404, "No such backup")
    return FileResponse(os.path.join(config.BACKUP_DIR, name), media_type="application/zip", filename=name)


@app.get("/.well-known/assetlinks.json")
def assetlinks():
    """Digital Asset Links: tells Android that the Play Store app and this website belong together, so the app shows
    the site full screen without a browser bar."""
    if not config.ANDROID_CERT_SHA256:
        return JSONResponse([])
    return JSONResponse([{"relation": ["delegate_permission/common.handle_all_urls"],
                          "target": {"namespace": "android_app", "package_name": config.ANDROID_PACKAGE,
                                     "sha256_cert_fingerprints": config.ANDROID_CERT_SHA256}}])


@app.get("/privacy", response_class=HTMLResponse)
def privacy():
    """Privacy policy for the employee app (the Play Store requires a public link)."""
    company = html.escape(config.COMPANY_NAME or "The company")
    contact = html.escape(config.PRIVACY_CONTACT or "your HR department")
    return open("app/static/privacy.html", encoding="utf-8").read().replace("{{COMPANY}}", company).replace("{{CONTACT}}", contact)


app.mount("/m", StaticFiles(directory="app/static/m", html=True), name="mobile")  # employee PWA (public shell)
app.mount("/static", StaticFiles(directory="app/static"), name="static")
