"""Excel helpers shared by the reports and the weekly e-mail: one coloured table per worksheet."""
from datetime import date

from . import reports

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TIMECARD_COLUMNS = [("name", "Employee Name"), ("device_user_id", "Employee ID"), ("department", "Department"),
                    ("date", "Date"), ("weekday", "Day"), ("day_type", "Day Type"),
                    ("scheduled_in", "Scheduled Login"), ("in", "Actual Login"),
                    ("scheduled_out", "Scheduled Logout"), ("out", "Actual Logout"),
                    ("late_min", "Late Arrival (min)"), ("early_in_min", "Early Login (min)"),
                    ("early_min", "Early Logout (min)"), ("hours", "Total Working Hours"), ("status", "Attendance Status"),
                    ("leave_type", "Leave Type"), ("week_off_duty", "Week-off duty"), ("note", "Note")]


def timecard_human(r: dict) -> dict:
    """A time card row with Excel headers, a real date, and 'Yes (rostered)' for approved week-off work."""
    return {label: (date.fromisoformat(r[k]) if k == "date" else ("Yes (rostered)" if r[k] else "") if k == "week_off_duty"
                    else r[k]) for k, label in TIMECARD_COLUMNS}


def safe(v):
    """Stop spreadsheet formula injection: a text cell starting with = + - @ would be executed by Excel."""
    return "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v


def xlsx_sheet(ws, title: str, rows: list[dict], header_rows: list[list] | None = None):
    """One table on a worksheet, with the leave colours: a leave day's row (leave_type) and IL / AL codes are filled
    yellow (Informed) or red (Adhoc), as are non-zero Informed / Adhoc leave counts; a legend goes under the table."""
    from openpyxl.styles import Font, PatternFill
    fills = {k: PatternFill("solid", fgColor=c) for k, c in reports.XLSX_FILLS.items()}
    ws.title = title[:31]
    for h in header_rows or []:      # e.g. the employee's name and ID above their own table
        ws.append(h)
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    if header_rows:
        ws.append([])
    top = ws.max_row + 1 if header_rows else 1
    if not rows:
        return
    keys = list(rows[0].keys())
    ws.append(keys)
    for cell in ws[top]:
        cell.font = Font(bold=True)
    kind_cols = {i + 1: k for i, h in enumerate(keys) for k, names in reports.LEAVE_COLUMNS.items() if h in names}
    lt = next((h for h in keys if h in ("leave_type", "Leave Type")), None)
    for r in rows:
        ws.append([safe(v) for v in r.values()])
        row = ws.max_row
        for cell in ws[row]:
            if isinstance(cell.value, date):
                cell.number_format = "DD-MM-YYYY"
        kind = r.get(lt) if lt else None
        if kind in fills:
            for cell in ws[row]:
                cell.fill = fills[kind]
        elif r.get("Week-off duty") or r.get("week_off_duty"):
            for cell in ws[row]:
                cell.fill = fills["WeekOff"]
        for cell in ws[row]:
            v = cell.value
            if v in ("IL", "AL"):
                cell.fill = fills["Informed" if v == "IL" else "Adhoc"]
            elif cell.column in kind_cols and isinstance(v, (int, float)) and v > 0:
                cell.fill = fills[kind_cols[cell.column]]
    for col in ws.columns:  # readable column widths
        ws.column_dimensions[col[0].column_letter].width = min(40, max(len(str(c.value or "")) for c in col[top - 1:]) + 2)
    ws.freeze_panes = ws.cell(top + 1, 1)
    ws.append([])
    for key, label in reports.XLSX_LEGEND:
        ws.append([label])
        ws.cell(ws.max_row, 1).fill = fills[key]


def inout_sheet(ws, title: str, rows: list[dict], start: str, end: str):
    """The time card as a grid: one row per employee, two columns (In, Out) per date. Leave days show IL / AL in
    yellow / red, weekly offs WO, holidays H and absent days A across both columns; a late check-in is in red, an
    approved (rostered) week-off duty is blue, and a missing punch stays empty (nothing is invented)."""
    from datetime import timedelta
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    fills = {k: PatternFill("solid", fgColor=c) for k, c in reports.XLSX_FILLS.items()}
    grey, sky = PatternFill("solid", fgColor="F1F5F9"), PatternFill("solid", fgColor="E0F2FE")
    thin = Side(style="thin", color="CBD5E1")
    box, mid = Border(left=thin, right=thin, top=thin, bottom=thin), Alignment(horizontal="center", vertical="center")
    ws.title = title[:31]
    days, d = [], date.fromisoformat(start)
    while d <= date.fromisoformat(end):
        days.append(d)
        d += timedelta(days=1)
    ws.cell(1, 1, "Date →").font = Font(bold=True)
    ws.cell(2, 1, "Employee").font = Font(bold=True)
    ws.cell(2, 2, "ID").font = Font(bold=True)
    for i, day in enumerate(days):
        c = 3 + 2 * i
        ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + 1)
        h = ws.cell(1, c, f"{day:%a %d/%m}")
        h.font, h.alignment = Font(bold=True), mid
        for j, label in enumerate(("In", "Out")):
            x = ws.cell(2, c + j, label)
            x.font, x.alignment, x.border = Font(bold=True), mid, box
        ws.column_dimensions[get_column_letter(c)].width = ws.column_dimensions[get_column_letter(c + 1)].width = 7
    by = {(r["device_user_id"], r["date"]): r for r in rows}
    people = sorted({(r["name"], r["device_user_id"]) for r in rows}, key=lambda p: (p[0].lower(), p[1]))
    codes = {"Weekly Off": ("WO", grey), "Holiday": ("H", sky), "Absent": ("A", None)}
    for n, (name, uid) in enumerate(people):
        row = 3 + n
        ws.cell(row, 1, safe(name)).font = Font(bold=True)
        ws.cell(row, 2, safe(uid))
        for i, day in enumerate(days):
            c, r = 3 + 2 * i, by.get((uid, day.isoformat()))
            cin, cout = ws.cell(row, c), ws.cell(row, c + 1)
            cin.border = cout.border = box
            cin.alignment = cout.alignment = mid
            if not r or not r["status"]:
                continue
            kind = r.get("leave_type")
            if r["status"] == "Leave" or r["status"] in codes:
                ws.merge_cells(start_row=row, start_column=c, end_row=row, end_column=c + 1)
                if r["status"] == "Leave":
                    cin.value, fill = ("IL" if kind == "Informed" else "AL" if kind == "Adhoc" else "LV"), fills.get(kind)
                else:
                    cin.value, fill = codes[r["status"]]
                if r.get("absent_adhoc"):          # absent without approved leave: counted as Adhoc
                    fill = fills["Adhoc"]
                if fill:
                    cin.fill = fill
                if r["status"] == "Absent":
                    cin.font = Font(bold=True, color="B91C1C")
                continue
            cin.value, cout.value = r["in"] or None, r["out"] or None
            if r["late_min"] > 0:
                cin.font = Font(bold=True, color="B91C1C")
            fill = fills.get(kind) if kind else fills["WeekOff"] if r.get("week_off_duty") else None
            if fill:
                cin.fill = cout.fill = fill
    ws.column_dimensions["A"].width = max([len("Employee")] + [len(p[0]) for p in people]) + 2
    ws.column_dimensions["B"].width = max([4] + [len(p[1]) for p in people]) + 2
    ws.freeze_panes = "C3"
    r0 = ws.max_row + 2
    legend = [("IL", fills["Informed"], "Informed leave"), ("AL", fills["Adhoc"], "Adhoc leave"), ("WO", grey, "Weekly off"),
              ("H", sky, "Holiday"), ("A", None, "Absent"), ("A ", fills["Adhoc"], "Absent without approved leave: counted as Adhoc"), ("10:25", None, "Late check-in (red)"),
              ("", fills["WeekOff"], "Rostered work on a weekly off")]
    for k, (code, fill, label) in enumerate(legend):
        x = ws.cell(r0 + k, 1, code or " ")
        if fill:
            x.fill = fill
        if code.strip() in ("A", "10:25"):
            x.font = Font(bold=True, color="B91C1C")
        ws.cell(r0 + k, 2, label)


def _hhmm(minutes: int) -> str:
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


# the summary line under each employee: (label, key, number of day columns it spans); 10 + 6 x 2 + 3 + 3 x 2 = 31
PERF_SUMMARY = [("Name", "name", 10), ("Present", "present", 2), ("HL", "holidays", 2), ("WO", "weekly_offs", 2),
                ("Absent", "absent", 2), ("Leave", "leave", 2), ("PaidDays", "paid_days", 2),
                ("Early Going", "early_min", 3), ("LateHrs.", "late_min", 2), ("WorkHrs", "work_min", 2),
                ("OvTim", "ot_min", 2)]
PERF_ROWS = [("Arrived Time", "in"), ("Dept. Time", "out"), ("Working Hrs.", "hours_min"), ("O.Time Hrs.", "ot_min"),
             ("Status", "code"), ("Shift", "shift")]


def performance_workbook(people: list[dict], start: str, end: str, company: str = "") -> bytes:
    """The old machine software's 'Monthly Performance Report', one sheet per branch: for each employee a department /
    designation line, a summary line and a day-by-day grid (arrival, departure, working hours, overtime, status,
    shift). Leave days are yellow (Informed) or red (Adhoc), as everywhere else; printed landscape on A4."""
    import io
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as col

    d1, d2 = date.fromisoformat(start), date.fromisoformat(end)
    days = [date.fromordinal(o) for o in range(d1.toordinal(), d2.toordinal() + 1)]
    ncols = 1 + max(len(days), 31)
    thin = Side(style="thin", color="000000")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    f = lambda size=7, bold=False, color="000000": Font(name="Verdana", size=size, bold=bold, color=color)
    fills = {k: PatternFill("solid", fgColor=c) for k, c in reports.XLSX_FILLS.items()}
    center, left = Alignment(horizontal="center", vertical="center"), Alignment(horizontal="left", vertical="center")
    month = d1.strftime("%b %Y") if (d1.year, d1.month) == (d2.year, d2.month) else f"{d1:%b %Y} - {d2:%b %Y}"

    def put(ws, r, c, v, span=1, font=None, align=center, border=True, fill=None):
        if span > 1:
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
        cell = ws.cell(r, c, safe(v))
        cell.font, cell.alignment = font or f(), align
        if fill:
            cell.fill = fill
        if border:
            for k in range(c, c + span):
                ws.cell(r, k).border = box
        return cell

    wb = Workbook()
    wb.remove(wb.active)
    branches: dict[str, list[dict]] = {}
    for p in people:
        branches.setdefault(p["branch"] or "No branch", []).append(p)
    used: set[str] = set()
    for branch, staff in (branches or {"No branch": []}).items():
        title = "".join(ch for ch in branch if ch not in '[]:*?/\\')[:31] or "Sheet"
        while title.lower() in used:
            title = title[:28] + f" {len(used)}"
        used.add(title.lower())
        ws = wb.create_sheet(title)
        ws.column_dimensions["A"].width = 13
        for c in range(2, ncols + 1):
            ws.column_dimensions[col(c)].width = 6
        third = (ncols - 1) // 3
        put(ws, 1, 1, f"Report from : {d1:%d-%m-%Y} To : {d2:%d-%m-%Y}", 1 + third, f(8), left, False)
        put(ws, 1, 2 + third, "Monthly Performance Report", third, Font(name="Arial", size=10, bold=True), center, False)
        put(ws, 1, 2 + 2 * third, f"Print Date {datetime.now():%d/%m/%Y}", ncols - 1 - 2 * third, f(8),
            Alignment(horizontal="right"), False)
        r = 3
        if company:
            put(ws, r, 1, f"Company Name : {company}", ncols, Font(name="Arial", size=8, bold=True, color="0000FF"), left, False)
            r += 1
        put(ws, r, 1, f"Branch : {branch}", ncols, Font(name="Arial", size=8, bold=True, color="0000FF"), left, False)
        r += 2
        if not staff:
            put(ws, r, 1, "No employees in this report", ncols, f(8), left, False)
        for p in staff:
            put(ws, r, 1, f"Dept:- {p['department']}", 11, f(7, True), left)
            put(ws, r, 12, f"Desig : {p['designation']}", 11, f(7, True), left)
            put(ws, r, 23, f"Attendance Month Of:- {month}", ncols - 22, f(7, True), left)
            r += 1
            put(ws, r, 1, "EmpCode", 1, f(8, True), left)
            c = 2
            for label, _, span in PERF_SUMMARY:
                put(ws, r, c, label, span, f(8, True), left if label == "Name" else center)
                c += span
            r += 1
            put(ws, r, 1, p["device_user_id"], 1, f(8), left)
            c = 2
            for _, key, span in PERF_SUMMARY:
                v = p[key]
                v = _hhmm(v) if key.endswith("_min") else (int(v) if isinstance(v, float) and v.is_integer() else v)
                put(ws, r, c, v, span, f(8), left if key == "name" else center)
                c += span
            r += 1
            for i in range(ncols - 1):       # empty boxes after the last date, as in the old report
                put(ws, r, 2 + i, f"{days[i].day:02d}" if i < len(days) else "", 1, f(7, True))
            put(ws, r, 1, "", 1)
            r += 1
            for label, key in PERF_ROWS:
                put(ws, r, 1, label, 1, f(7), left)
                for i in range(ncols - 1):
                    x = p["days"].get(days[i].isoformat()) if i < len(days) else None
                    v = "" if x is None else _hhmm(x[key]) if key in ("hours_min", "ot_min") else x[key]
                    fill = None
                    if x and key == "code":
                        fill = fills.get(x["leave_type"]) or (fills["Adhoc"] if x["absent_adhoc"] else
                                                              fills["WeekOff"] if x["week_off_duty"] else None)
                    cell = put(ws, r, 2 + i, v, 1, f(5 if key == "shift" else 6), center, True, fill)
                    if key == "shift":
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if key == "shift":
                    ws.row_dimensions[r].height = 24
                r += 1
            r += 1
        r += 1
        for key, text in reports.XLSX_LEGEND:
            put(ws, r, 1, "", 1, fill=fills[key])
            put(ws, r, 2, text, 10, f(7), left, False)
            r += 1
        put(ws, r, 1, "Codes: P present, P-LT late, P/2 half day, POW worked on a weekly off, A absent, WO weekly off, "
                      "HL holiday, IL Informed leave, AL Adhoc leave", ncols, f(7), left, False)
        ps = ws.page_setup
        ps.orientation, ps.paperSize, ps.fitToWidth, ps.fitToHeight = "landscape", ws.PAPERSIZE_A4, 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.3
        ws.oddFooter.left.text = "Page No &P of &N"
        ws.oddFooter.right.text = "Approved BY :________________"
        ws.sheet_view.showGridLines = False
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
