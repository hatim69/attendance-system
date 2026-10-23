"""Weekly attendance report: every Monday morning HR gets an Excel file of the month so far.

It covers the 1st of the month up to the Sunday just ended. On the first Monday of a month (when that Sunday is in
the first 6 days of the month) it covers the whole previous month instead, as its closing report, so every month
gets one complete report. Sheets: Summary (one row per employee), Late arrivals, Leave, and Time Card (every
employee, day by day), with the yellow / red leave colours."""
import asyncio
import io
import json
import logging
from datetime import date, datetime, timedelta

from . import config, db, lates, mailer, reports
from .timeutil import now_local
from .xlsx import timecard_human, xlsx_sheet

log = logging.getLogger("weekly")


def period(today: date) -> tuple[date, date]:
    end = today - timedelta(days=(today.weekday() + 1) % 7 or 7)    # the last Sunday before today
    if end.day <= 6:                                                 # a new month has just started: close the last one
        last = end.replace(day=1) - timedelta(days=1)
        return last.replace(day=1), last
    return end.replace(day=1), end


def summary(rows: list[dict]) -> list[dict]:
    out: dict[str, dict] = {}
    for r in rows:
        x = out.setdefault(r["device_user_id"], {
            "Employee Name": r["name"], "Employee ID": r["device_user_id"], "Department": r["department"],
            "Present": 0, "Late": 0, "Half Day": 0, "Absent": 0, "Informed Leave": 0, "Adhoc Leave": 0, "Holidays": 0,
            "Weekly Offs": 0, "Early Logins": 0, "Early Logouts": 0, "Missing Punches": 0, "Hours": 0.0})
        st = r["status"]
        if st in ("Present", "Late", "Half Day", "Absent"):
            x[st] += 1
        x["Holidays"] += st == "Holiday"
        x["Weekly Offs"] += st == "Weekly Off"
        if r["leave_type"] in ("Informed", "Adhoc"):
            x[r["leave_type"] + " Leave"] += 0.5 if st == "Half Day" else 1
        x["Early Logins"] += r["early_in_min"] > 0
        x["Early Logouts"] += r["early_min"] > 0
        x["Missing Punches"] += "missing" in r["note"]
        x["Hours"] = round(x["Hours"] + r["hours"], 2)
    return sorted(out.values(), key=lambda x: (x["Employee Name"].lower(), x["Employee ID"]))


def build(a: date, b: date) -> tuple[bytes, dict]:
    """The workbook for a..b and a few totals for the e-mail text."""
    from openpyxl import Workbook
    rows = reports.timecards(a.isoformat(), b.isoformat())
    summ = summary(rows)
    lt = lates.counts(b)
    label = {"ok": "", "limit": "Limit reached", "over": "Over the limit: HR review"}
    late_rows = [{"Employee Name": x["name"], "Employee ID": x["device_user_id"], "Department": x["department"],
                  "Late Arrivals": x["count"], "Limit": x["limit"], "Flag": label[x["state"]],
                  "Dates (minutes late)": ", ".join(f"{l['date'][8:]}/{l['date'][5:7]} ({l['late_min']})" for l in x["lates"])}
                 for x in lt if x["count"]]
    leave_rows = [{"Employee Name": r["name"], "Employee ID": r["device_user_id"], "From": date.fromisoformat(r["start_date"]),
                   "To": date.fromisoformat(r["end_date"]), "Leave Type": r["type"], "Half Day": "Yes" if r["half"] else "",
                   "Status": r["status"], "Pay Cut (x day's salary)": r["lop"] or "", "Note": r["note"]}
                  for r in db.conn().execute(
                      "SELECT l.*, COALESCE(e.name, l.device_user_id) name FROM leaves l LEFT JOIN employees e "
                      "ON e.device_user_id=l.device_user_id WHERE l.status IN ('approved','pending') "
                      "AND l.start_date<=? AND l.end_date>=? ORDER BY l.start_date, name", (b.isoformat(), a.isoformat()))]
    wb = Workbook()
    xlsx_sheet(wb.active, "Summary", summ or [{"Employee Name": "No attendance in this period"}],
               header_rows=[["Attendance report", f"{a:%d %b %Y} to {b:%d %b %Y}"],
                            ["Late arrivals counted per", lates.policy()["period_label"]]])
    xlsx_sheet(wb.create_sheet(), "Late arrivals", late_rows or [{"Employee Name": "No late arrivals"}])
    xlsx_sheet(wb.create_sheet(), "Leave", leave_rows or [{"Employee Name": "No leave in this period"}])
    xlsx_sheet(wb.create_sheet(), "Time Card", [timecard_human(r) for r in rows] or [{"Employee Name": "No records"}])
    out = io.BytesIO()
    wb.save(out)
    totals = {"employees": len(summ), "late": sum(x["Late"] for x in summ), "absent": sum(x["Absent"] for x in summ),
              "informed": sum(x["Informed Leave"] for x in summ), "adhoc": sum(x["Adhoc Leave"] for x in summ),
              "flagged": [x["name"] for x in lt if x["state"] != "ok"]}
    return out.getvalue(), totals


def status() -> dict:
    row = db.conn().execute("SELECT value FROM settings WHERE key='weekly_report_last'").fetchone()
    return json.loads(row["value"]) if row else {}


def run(now: datetime | None = None) -> dict:
    """Build and e-mail the report for the period that ends on the last Sunday; remember how it went."""
    now = now or now_local()
    a, b = period(now.date())
    st = {"week": now.date().isoformat(), "at": now.strftime("%Y-%m-%d %H:%M:%S"), "from": a.isoformat(),
          "to": b.isoformat(), "sent_to": "", "error": ""}
    try:
        data, t = build(a, b)
        if not mailer.configured():
            raise RuntimeError("e-mail is not set up (SMTP_HOST in .env)")
        body = (f"Attendance report for {a:%d %b %Y} to {b:%d %b %Y} ({t['employees']} employees).\n\n"
                f"Late arrivals: {t['late']}\nAbsent days: {t['absent']}\n"
                f"Informed leave (yellow): {t['informed']:g} days\nAdhoc leave (red): {t['adhoc']:g} days\n"
                + (f"\nAt or over the late-arrival limit (HR review): {', '.join(t['flagged'])}\n" if t["flagged"] else "")
                + "\nThe attached Excel file has the summary, late arrivals, leave and the full time card.")
        name = f"attendance-report-{a.isoformat()}-to-{b.isoformat()}.xlsx"
        mailer.send([config.HR_EMAIL], f"Attendance report {a:%d %b} - {b:%d %b %Y}", body, (name, data))
        st["sent_to"] = config.HR_EMAIL
    except Exception as e:
        log.exception("weekly report failed")
        st["error"] = f"{type(e).__name__}: {e}"[:300]
    c = db.conn()
    c.execute("INSERT INTO settings(key,value) VALUES ('weekly_report_last',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
              (json.dumps(st),))
    c.commit()
    return st


def due(now: datetime) -> bool:
    return now.weekday() == 0 and now.hour >= config.WEEKLY_REPORT_HOUR and status().get("week") != now.date().isoformat()


async def loop():
    """Monday from WEEKLY_REPORT_HOUR (or as soon after as the server is running), once a week."""
    while True:
        try:
            now = now_local()
            if due(now):
                await asyncio.to_thread(run, now)
        except Exception:
            log.exception("weekly report check failed")
        await asyncio.sleep(600)

